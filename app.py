"""
MatrixDev · Gestión y arquitectura de proyectos de software (Tesis)

Ejecutar con:  streamlit run MATRIXDEVTESIS.py
Requiere Streamlit 1.40 o superior.
Dependencias: streamlit, pandas, numpy, plotly, streamlit-mermaid, openpyxl, python-docx
"""
import io
import json
import zipfile
from datetime import date

import numpy as np
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st
from streamlit_mermaid import st_mermaid

import openpyxl
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from docx import Document
from docx.shared import Pt, RGBColor

# =====================================================================
# 1. CONFIGURACIÓN GENERAL
# =====================================================================
st.set_page_config(page_title="MatrixDev · Tesis", layout="wide", page_icon=":material/hub:",
                   initial_sidebar_state="expanded")


def _version_streamlit():
    try:
        return tuple(int(p) for p in st.__version__.split(".")[:2])
    except Exception:
        return (1, 0)


# Streamlit >= 1.50 reemplazó use_container_width por width="stretch"
ANCHO = {"width": "stretch"} if _version_streamlit() >= (1, 50) else {"use_container_width": True}

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# =====================================================================
# 2. ESTADO DE LA APLICACIÓN
# =====================================================================
PLANTILLAS_FLUJO = {
    "Autenticación": """graph TD
    A[Usuario] -->|Credenciales| B[API Login]
    B -->|Validar| C{"¿Es Válido?"}
    C -->|Sí| D[Generar Token JWT]
    C -->|No| E[Error 401 Unauthorized]""",

    "Procesamiento API": """graph LR
    Cliente[Cliente] -->|Request POST| Router[Router]
    Router --> Controller[Controller]
    Controller -->|Query| DB[(Base de Datos)]
    DB -->|Respuesta| Controller
    Controller -->|JSON| Cliente""",

    "Secuencia de Usuario": """sequenceDiagram
    autonumber
    actor Cliente
    participant Servidor
    participant BD as Base de Datos

    Cliente->>Servidor: POST /login
    Servidor->>BD: Consulta Usuario
    BD-->>Servidor: Datos OK
    Servidor-->>Cliente: 200 OK + Token""",

    "Crear desde cero": """graph TD
    A[Inicio] --> B[Tu Nuevo Proceso]
    B --> C{"¿Aprobado?"}
    C -->|Sí| D[Resultado Éxito]
    C -->|No| E[Resultado Fallo]"""
}

# Datos del proyecto (no son widgets)
# El proyecto parte vacío: el inicio se va completando a medida que se ingresan datos
DATOS_INICIALES = {
    "rf_list": [], "rnf_list": [], "kanban_tasks": [], "edt_list": [], "riesgos_list": [],
    "mc_semilla_val": 42,
}

# Datos de muestra que se cargan con el botón "Cargar ejemplo"
DATOS_EJEMPLO = {
    "rf_list": ["Autenticación OAuth2.", "CRUD de usuarios."],
    "rnf_list": ["Latencia < 200ms.", "Cifrado AES-256 en base de datos."],
    "kanban_tasks": [
        {"Tarea": "Modelo BD", "Estado": "Completado"},
        {"Tarea": "Endpoints API", "Estado": "En Proceso"},
    ],
    "edt_list": [
        {"Fase": "Fase 1: Preparación", "Subfase": "Limpieza y Despeje", "Paquete": "Despejar Habitación y Muebles",
         "Horas": 5, "Costo": 20000, "Documentos": "Plan de Trabajo"},
        {"Fase": "Fase 1: Preparación", "Subfase": "Pintura Aerosol", "Paquete": "Pintar Arbusto / Detalles",
         "Horas": 4, "Costo": 35000, "Documentos": "Ficha Técnica Aerosol"},
        {"Fase": "Fase 2: Aplicación Manual", "Subfase": "Capas de Pintura", "Paquete": "Mano de Pintura Muros (1 y 2)",
         "Horas": 16, "Costo": 80000, "Documentos": "Guía de Pintado Manual"},
        {"Fase": "Fase 2: Aplicación Manual", "Subfase": "Documentación y Entrega",
         "Paquete": "Informe de Inspección y Cotizaciones", "Horas": 6, "Costo": 15000,
         "Documentos": "Cotización + Certificados"},
    ],
    "riesgos_list": [
        {"Riesgo": "Retraso en entrega API", "Probabilidad": 3, "Impacto": 4},
        {"Riesgo": "Falla en servidor cloud", "Probabilidad": 2, "Impacto": 5},
        {"Riesgo": "Incompatibilidad de navegador", "Probabilidad": 1, "Impacto": 2},
    ],
    "objetivo_text": "Automatizar el flujo de inventario con una arquitectura minimalista.",
    "integrantes": "Juan Pérez",
}

# Valores de widgets que deben conservarse al cambiar de página
WIDGETS_PERSISTENTES = {
    "tema": "claro",
    "nombre_proj": "MatrixDev Core",
    "integrantes": "",
    "objetivo_text": "",
    "modelo_mermaid": "graph TD\n    A[Inicio] --> B{Validar}\n    B -->|Sí| C[Éxito]\n    B -->|No| D[Error]",
    "select_tipo_flujo": "Autenticación",
    "flujo_codigo": PLANTILLAS_FLUJO["Autenticación"],
    "rad_edt_metric": "Horas",
    # Monte Carlo
    "mc_fuente": "Estimación manual (3 puntos)",
    "mc_dist_type": "Beta-PERT",
    "mc_horas_dia": 8.0, "mc_pct_opt": 20, "mc_pct_pes": 60,
    "mc_opt": 10.0, "mc_prob": 20.0, "mc_pes": 45.0,
    "mc_mean": 22.0, "mc_std": 5.0,
    "mc_min_u": 12.0, "mc_max_u": 40.0,
    "mc_cfijo": 100000, "mc_cdiario": 15000, "mc_var_costo": 5,
    "mc_dias_meta": 25.0, "mc_costo_meta": 500000, "mc_target_perc": 85,
    "mc_f_inicio": date.today(), "mc_habiles": True,
    "mc_sims": 10000, "mc_seed": True,
}

for _k, _v in DATOS_INICIALES.items():
    if _k not in st.session_state:
        st.session_state[_k] = _v
for _k, _v in WIDGETS_PERSISTENTES.items():
    if _k not in st.session_state or st.session_state[_k] is None:
        st.session_state[_k] = _v
    else:
        # Reasignar evita que Streamlit borre el valor cuando el widget no está en la página actual
        st.session_state[_k] = st.session_state[_k]

ESTADOS_KANBAN = ["Pendiente", "En Proceso", "Completado"]

# =====================================================================
# 3. TEMAS
#   claro       → celeste y blanco
#   oscuro      → negro y blanco
#   informatico → negro con rojo neón
# =====================================================================
TEMAS = {
    "claro": {
        "nombre": "Claro", "base": "light", "icono": ":material/light_mode:",
        "bg": "#f5fafd", "card": "#ffffff", "text": "#0d2433", "sub": "#5b7385",
        "accent": "#0b8fd6", "accent2": "#5cc3f0", "soft": "#e4f3fc", "border": "#dcebf4", "grid": "#e9f3f9",
        "input": "#ffffff", "ok": "#0f9d6e", "warn": "#d88a07", "bad": "#d9473f", "on_accent": "#ffffff",
        "fuente": "'IBM Plex Sans', 'Segoe UI', Helvetica, Arial, sans-serif",
        "titulos": "'IBM Plex Sans', 'Segoe UI', Helvetica, Arial, sans-serif",
        "neon": False,
        "fondo": "radial-gradient(1100px 380px at 100% -5%, rgba(11,143,214,0.10), transparent 65%)",
    },
    "oscuro": {
        "nombre": "Oscuro", "base": "dark", "icono": ":material/dark_mode:",
        "bg": "#000000", "card": "#0a0a0a", "text": "#fafafa", "sub": "#9b9b9b",
        "accent": "#ffffff", "accent2": "#a3a3a3", "soft": "#161616", "border": "#232323", "grid": "#171717",
        "input": "#101010", "ok": "#4ade80", "warn": "#facc15", "bad": "#f87171", "on_accent": "#000000",
        "fuente": "'IBM Plex Sans', 'Segoe UI', Helvetica, Arial, sans-serif",
        "titulos": "'IBM Plex Sans', 'Segoe UI', Helvetica, Arial, sans-serif",
        "neon": False,
        "fondo": "linear-gradient(transparent, transparent)",
    },
    "informatico": {
        "nombre": "Informático", "base": "dark", "icono": ":material/terminal:",
        "bg": "#050506", "card": "#0b090a", "text": "#f2f2f2", "sub": "#958b8d",
        "accent": "#ff2a4f", "accent2": "#ff7a91", "soft": "#1a0a0e", "border": "#2b1118", "grid": "#170a0d",
        "input": "#0e0a0b", "ok": "#2ee88a", "warn": "#ffb020", "bad": "#ff4040", "on_accent": "#050506",
        "fuente": "'IBM Plex Sans', 'Segoe UI', Helvetica, Arial, sans-serif",
        "titulos": "'IBM Plex Mono', Consolas, monospace",
        "neon": True,
        "fondo": "linear-gradient(rgba(255,42,79,0.035) 1px, transparent 1px) 0 0 / 28px 28px, "
                 "linear-gradient(90deg, rgba(255,42,79,0.035) 1px, transparent 1px) 0 0 / 28px 28px",
    },
}

if st.session_state.get("tema") not in TEMAS:
    st.session_state.tema = st.session_state.get("_tema_prev", "claro")
st.session_state._tema_prev = st.session_state.tema
T = TEMAS[st.session_state.tema]
FUENTE = T["fuente"]


def rgba(hex_color, alpha):
    h = hex_color.lstrip("#")
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    return f"rgba({r},{g},{b},{alpha})"


def aplicar_tema_nativo(tema):
    """Sincroniza el tema nativo de Streamlit (sliders, toggles, menús, calendario) con el tema elegido."""
    cambio = False
    try:
        from streamlit import config as _cfg
    except Exception:
        return False
    opciones = {
        "theme.base": tema["base"],
        "theme.primaryColor": tema["accent"],
        "theme.backgroundColor": tema["bg"],
        "theme.secondaryBackgroundColor": tema["input"] if tema["base"] == "dark" else tema["soft"],
        "theme.textColor": tema["text"],
    }
    for clave, valor in opciones.items():
        try:
            if _cfg.get_option(clave) != valor:
                _cfg.set_option(clave, valor)
                cambio = True
        except Exception:
            pass
    return cambio



NEON = T["neon"]
GLOW = f"0 0 0 1px {rgba(T['accent'], 0.35)}, 0 0 22px {rgba(T['accent'], 0.28)}" if NEON else f"0 8px 24px {rgba(T['accent'], 0.10)}"
GLOW_TXT = f"0 0 14px {rgba(T['accent'], 0.55)}" if NEON else "none"

# =====================================================================
# 4. ESTILOS (CSS)
# =====================================================================
st.markdown(f"""
<style>
@import url('https://fonts.googleapis.com/css2?family=IBM+Plex+Sans:wght@400;500;600;700&family=IBM+Plex+Mono:wght@400;500;600&display=swap');

@keyframes mdRise {{ from {{ opacity: 0; transform: translateY(8px); }} to {{ opacity: 1; transform: none; }} }}
@keyframes mdFade {{ from {{ opacity: 0; }} to {{ opacity: 1; }} }}
@keyframes mdDraw {{ from {{ transform: scaleX(0); }} to {{ transform: scaleX(1); }} }}
@keyframes mdGrow {{ from {{ width: 0; }} }}
@keyframes mdDot {{ from {{ opacity: 0; transform: scale(.4); }} to {{ opacity: 1; transform: scale(1); }} }}

/* ---------- Base ---------- */
.stApp {{ background: {T['fondo']}, {T['bg']} !important; background-attachment: fixed !important; font-family: {FUENTE}; }}
[data-testid="stHeader"] {{ background: transparent !important; }}
#MainMenu, footer, [data-testid="stDecoration"] {{ display: none !important; }}
.block-container, [data-testid="stMainBlockContainer"] {{ padding-top: 2.2rem !important; max-width: 1320px; }}
.stMarkdown, .stMarkdown p, .stMarkdown li, label, input, textarea, button {{ font-family: {FUENTE} !important; }}
h1, h2, h3, h4 {{ font-family: {T['titulos']} !important; letter-spacing: -0.015em; color: {T['text']} !important; }}
::selection {{ background: {rgba(T['accent'], 0.25)}; }}
::-webkit-scrollbar {{ width: 9px; height: 9px; }}
::-webkit-scrollbar-thumb {{ background: {rgba(T['sub'], 0.35)}; border-radius: 9px; }}
.ico {{ display: inline-block; vertical-align: middle; flex: none; }}

/* ---------- Barra lateral ---------- */
[data-testid="stSidebar"] {{ background: {T['card']} !important; border-right: 1px solid {T['border']}; }}
[data-testid="stSidebarNav"] a, [data-testid="stSidebarNavLink"] {{ border-radius: 8px !important; transition: background-color .2s ease; }}
[data-testid="stSidebarNavLink"][aria-current="page"], [data-testid="stSidebarNav"] a[aria-current="page"] {{
  background: {rgba(T['accent'], 0.12)} !important; }}
[data-testid="stSidebarNavLink"][aria-current="page"] span {{ color: {T['accent']} !important; font-weight: 600; }}
[data-testid="stNavSectionHeader"], [data-testid="stSidebarNavSeparator"] {{ color: {T['sub']} !important; }}

/* ---------- Controles ---------- */
div[data-baseweb="input"], div[data-baseweb="select"] > div, div[data-baseweb="textarea"],
[data-testid="stTextInputRootElement"], [data-testid="stNumberInputContainer"], [data-testid="stTextAreaRootElement"] {{
  border-radius: 8px !important; transition: border-color .2s ease, box-shadow .2s ease; }}
[data-testid="stTextInputRootElement"]:focus-within, [data-testid="stNumberInputContainer"]:focus-within,
[data-testid="stTextAreaRootElement"]:focus-within {{ box-shadow: 0 0 0 3px {rgba(T['accent'], 0.16)}; }}
.stButton>button, .stDownloadButton>button, .stFormSubmitButton>button, [data-testid="stPageLink"] a {{
  border-radius: 8px !important; font-weight: 500;
  transition: border-color .2s ease, background-color .2s ease, box-shadow .25s ease, transform .15s ease; }}
.stButton>button:hover, .stDownloadButton>button:hover, .stFormSubmitButton>button:hover {{
  border-color: {T['accent']} !important; color: {T['accent']} !important; box-shadow: {GLOW if NEON else 'none'}; }}
.stButton>button:active, .stDownloadButton>button:active {{ transform: scale(.98); }}
[data-testid="stExpander"] details {{ border-radius: 12px !important; border-color: {T['border']} !important; background: {T['card']}; }}
[data-testid="stForm"] {{ border-radius: 12px !important; border-color: {T['border']} !important; background: {T['card']}; }}
[data-testid="stVerticalBlockBorderWrapper"], [data-testid="stLayoutWrapper"] > [data-testid="stVerticalBlock"] {{ border-color: {T['border']} !important; }}
[data-testid="stTab"] p, .stTabs button p {{ font-weight: 500; }}
[data-testid="stTabPanel"], .stTabs [data-baseweb="tab-panel"] {{ animation: mdFade .3s ease both; padding-top: 8px; }}
hr {{ border-color: {T['border']} !important; }}
iframe {{ border-radius: 10px; }}

/* Accesos a módulos (st.page_link) */
[data-testid="stPageLink"] a {{ border: 1px solid {T['border']}; background: {T['card']}; padding: 10px 14px !important; }}
[data-testid="stPageLink"] a:hover {{ border-color: {T['accent']}; background: {rgba(T['accent'], 0.06)}; transform: translateY(-1px); }}

/* ---------- Encabezado de página ---------- */
.mx-head {{ position: relative; display: flex; justify-content: space-between; align-items: flex-end; gap: 24px;
  flex-wrap: wrap; padding-bottom: 18px; margin-bottom: 24px; animation: mdRise .45s ease both; }}
.mx-head::after {{ content: ""; position: absolute; left: 0; right: 0; bottom: 0; height: 1px; background: {T['border']}; }}
.mx-head::before {{ content: ""; position: absolute; left: 0; bottom: 0; height: 2px; width: 120px; background: {T['accent']};
  transform-origin: left; animation: mdDraw .7s .15s cubic-bezier(.2,.7,.2,1) both; z-index: 1; box-shadow: {GLOW if NEON else 'none'}; }}
.mx-head .l {{ display: flex; gap: 16px; align-items: flex-start; }}
.mx-tile {{ width: 46px; height: 46px; border-radius: 12px; display: grid; place-items: center; flex: none;
  background: {T['soft']}; color: {T['accent']}; border: 1px solid {rgba(T['accent'], 0.22)}; box-shadow: {GLOW if NEON else 'none'}; }}
.mx-head h1 {{ font-size: 1.75rem; font-weight: 600; margin: 0 !important; padding: 0 !important; line-height: 1.2; text-shadow: {GLOW_TXT}; }}
.mx-head .d {{ color: {T['sub']}; font-size: .97rem; margin-top: 4px; max-width: 72ch; line-height: 1.5; }}
.mx-head .meta {{ text-align: right; font-size: .84rem; color: {T['sub']}; line-height: 1.6; }}
.mx-head .meta b {{ color: {T['text']}; font-weight: 600; }}

/* ---------- Títulos de sección ---------- */
.mx-sec {{ display: flex; align-items: flex-start; gap: 10px; margin: 8px 0 12px 0; animation: mdRise .4s ease both; }}
.mx-sec .ico {{ color: {T['accent']}; margin-top: 2px; }}
.mx-sec .t {{ font-family: {T['titulos']}; font-size: 1.08rem; font-weight: 600; color: {T['text']}; }}
.mx-sec .s {{ color: {T['sub']}; font-size: .9rem; margin-top: 2px; max-width: 80ch; }}

/* ---------- Franja de indicadores ---------- */
.mx-strip {{ display: grid; grid-template-columns: repeat(var(--n, 4), minmax(0, 1fr)); gap: 1px; background: {T['border']};
  border: 1px solid {T['border']}; border-radius: 14px; overflow: hidden; }}
.mx-strip .c {{ background: {T['card']}; padding: 16px 18px; animation: mdRise .45s ease both; animation-delay: var(--d, 0s);
  transition: background-color .2s ease; }}
.mx-strip .c:hover {{ background: {rgba(T['accent'], 0.04)}; }}
.mx-strip .lb {{ display: flex; align-items: center; gap: 7px; color: {T['sub']}; font-size: .8rem; font-weight: 500; }}
.mx-strip .v {{ font-size: 1.45rem; font-weight: 600; margin-top: 8px; color: {T['text']}; font-variant-numeric: tabular-nums;
  font-family: {T['titulos']}; }}
.mx-strip .c.hl .v {{ color: {T['accent']}; text-shadow: {GLOW_TXT}; }}
.mx-strip .sb {{ color: {T['sub']}; font-size: .8rem; margin-top: 3px; }}

/* ---------- Panel / tarjeta ---------- */
.mx-panel {{ background: {T['card']}; border: 1px solid {T['border']}; border-radius: 14px; padding: 18px 20px;
  animation: mdRise .45s ease both; animation-delay: var(--d, 0s); }}
.mx-panel .pt {{ display: flex; align-items: center; gap: 8px; font-weight: 600; margin-bottom: 10px; font-family: {T['titulos']}; }}
.mx-panel .pt .ico {{ color: {T['accent']}; }}
.mx-row {{ display: flex; justify-content: space-between; gap: 14px; padding: 8px 0; border-bottom: 1px solid {T['grid']}; font-size: .92rem; }}
.mx-row:last-child {{ border-bottom: none; }}
.mx-row .k {{ color: {T['sub']}; }}
.mx-row .v {{ font-weight: 600; text-align: right; font-variant-numeric: tabular-nums; }}

.pill {{ display: inline-flex; align-items: center; gap: 5px; padding: 2px 10px; border-radius: 999px; font-size: .78rem; font-weight: 600; }}
.pill.ok {{ color: {T['ok']}; background: {rgba(T['ok'], 0.12)}; }}
.pill.bad {{ color: {T['bad']}; background: {rgba(T['bad'], 0.12)}; }}
.pill.warn {{ color: {T['warn']}; background: {rgba(T['warn'], 0.14)}; }}
.pill.neu {{ color: {T['sub']}; background: {rgba(T['sub'], 0.12)}; }}

/* ---------- Veredicto ---------- */
.mx-verdict {{ position: relative; display: grid; grid-template-columns: 1fr auto; gap: 26px; align-items: center;
  background: {T['card']}; border: 1px solid {T['border']}; border-radius: 14px; padding: 22px 24px 22px 28px; overflow: hidden;
  animation: mdRise .5s ease both; box-shadow: {GLOW if NEON else 'none'}; }}
.mx-verdict::before {{ content: ""; position: absolute; left: 0; top: 0; bottom: 0; width: 4px; background: var(--c); }}
.mx-verdict .tag {{ display: inline-flex; align-items: center; gap: 7px; color: var(--c); font-weight: 600; font-size: .86rem; }}
.mx-verdict .big {{ font-family: {T['titulos']}; font-size: 1.3rem; font-weight: 600; line-height: 1.35; margin: 8px 0 8px 0; max-width: 40ch; }}
.mx-verdict .big b {{ color: var(--c); }}
.mx-verdict .txt {{ color: {T['sub']}; font-size: .93rem; line-height: 1.6; max-width: 62ch; }}
.mx-verdict .txt b {{ color: {T['text']}; }}
.mx-dots {{ display: grid; grid-template-columns: repeat(10, 11px); gap: 4px; }}
.mx-dots i {{ width: 11px; height: 11px; border-radius: 3px; animation: mdDot .3s ease both; }}
.mx-dots i.ok {{ background: {T['ok']}; }}
.mx-dots i.mid {{ background: {T['warn']}; }}
.mx-dots i.bad {{ background: {rgba(T['bad'], 0.8)}; }}
.mx-leg {{ display: flex; gap: 12px; font-size: .74rem; color: {T['sub']}; margin-top: 10px; flex-wrap: wrap; }}
.mx-leg span::before {{ content: ""; display: inline-block; width: 8px; height: 8px; border-radius: 2px; margin-right: 5px; background: var(--c); }}
@media (max-width: 760px) {{ .mx-verdict {{ grid-template-columns: 1fr; }} }}
@media (max-width: 900px) {{ .mx-strip {{ grid-template-columns: repeat(2, minmax(0, 1fr)); }} }}

/* ---------- Barras de probabilidad ---------- */
.mx-bar {{ padding: 9px 0; }}
.mx-bar .top {{ display: flex; justify-content: space-between; gap: 10px; font-size: .9rem; }}
.mx-bar .top .k {{ color: {T['sub']}; }}
.mx-bar .top .v {{ font-weight: 600; color: var(--c); font-variant-numeric: tabular-nums; }}
.mx-bar .tr {{ position: relative; height: 8px; border-radius: 8px; background: {rgba(T['sub'], 0.14)}; margin-top: 7px; }}
.mx-bar .fl {{ height: 100%; border-radius: 8px; background: var(--c); width: var(--w); animation: mdGrow 1s cubic-bezier(.2,.7,.2,1) both;
  box-shadow: {('0 0 10px ' + rgba(T['accent'], 0.5)) if NEON else 'none'}; }}
.mx-bar .mk {{ position: absolute; top: -4px; bottom: -4px; width: 2px; border-radius: 2px; background: {T['text']}; left: var(--m); }}

/* ---------- Respuesta destacada ---------- */
.mx-answer {{ background: {T['soft']}; border: 1px solid {rgba(T['accent'], 0.22)}; border-radius: 10px; padding: 12px 15px;
  margin-top: 4px; animation: mdFade .3s ease both; font-size: .95rem; }}
.mx-answer b {{ color: {T['accent']}; font-size: 1.1rem; }}

/* ---------- Kanban ---------- */
.mx-kb {{ background: {T['card']}; border: 1px solid {T['border']}; border-radius: 14px; padding: 14px; min-height: 150px;
  animation: mdRise .45s ease both; animation-delay: var(--d, 0s); }}
.mx-kb .h {{ display: flex; justify-content: space-between; align-items: center; font-weight: 600; margin-bottom: 12px; }}
.mx-kb .h span.l {{ display: flex; align-items: center; gap: 8px; color: var(--c); }}
.mx-kb .h span.n {{ color: {T['sub']}; font-size: .8rem; font-weight: 600; border: 1px solid {T['border']}; border-radius: 999px; padding: 0 9px; }}
.mx-kb .it {{ border: 1px solid {T['border']}; border-radius: 8px; padding: 9px 11px; margin-bottom: 8px; font-size: .9rem;
  background: {T['bg']}; transition: border-color .2s ease, transform .2s ease; }}
.mx-kb .it:hover {{ border-color: var(--c); transform: translateX(2px); }}
.mx-kb .empty {{ color: {T['sub']}; font-size: .85rem; }}

/* ---------- Lista de recomendaciones ---------- */
.mx-todo {{ display: flex; gap: 12px; align-items: flex-start; padding: 10px 0; border-bottom: 1px solid {T['grid']}; font-size: .93rem; }}
.mx-todo:last-child {{ border-bottom: none; }}
.mx-todo .ico {{ margin-top: 1px; }}

/* ---------- Tablas ---------- */
.mx-tw {{ overflow-x: auto; border: 1px solid {T['border']}; border-radius: 12px; background: {T['card']}; animation: mdRise .45s ease both; }}
.mx-table {{ width: 100%; border-collapse: collapse; font-size: .9rem; margin: 0 !important; }}
.mx-table th, .mx-table td {{ border: none !important; }}
.mx-table th {{ text-align: left; padding: 11px 14px; color: {T['sub']}; font-weight: 600; font-size: .8rem; background: {T['soft']};
  border-bottom: 1px solid {T['border']} !important; white-space: nowrap; }}
.mx-table td {{ padding: 10px 14px; border-bottom: 1px solid {T['grid']} !important; font-variant-numeric: tabular-nums; }}
.mx-table tr:last-child td {{ border-bottom: none !important; }}
.mx-table tbody tr {{ transition: background-color .15s ease; }}
.mx-table tbody tr:hover {{ background: {rgba(T['accent'], 0.05)}; }}
.mx-table tr.hl td {{ background: {rgba(T['accent'], 0.10)}; font-weight: 600; }}
.mx-table tr.hl td:first-child {{ box-shadow: inset 3px 0 0 {T['accent']}; }}

/* ---------- Pasos del formulario Monte Carlo ---------- */
.mx-step {{ display: flex; align-items: center; gap: 9px; font-weight: 600; font-size: .95rem; margin: 6px 0 8px 0;
  padding-top: 6px; font-family: {T['titulos']}; }}
.mx-step .n {{ width: 22px; height: 22px; border-radius: 6px; display: inline-grid; place-items: center; font-size: .75rem;
  background: {T['accent']}; color: {T['on_accent']}; box-shadow: {GLOW if NEON else 'none'}; }}

/* Marca de la barra lateral */
.mx-brand {{ display: flex; align-items: center; gap: 10px; font-family: {T['titulos']}; font-weight: 700; font-size: 1.05rem; margin: 2px 0 4px 0; }}
.mx-brand .ico {{ color: {T['accent']}; }}
.mx-brand small {{ display: block; font-weight: 400; font-size: .78rem; color: {T['sub']}; font-family: {FUENTE}; }}

/* ---------- Navegación lateral más visible ---------- */
[data-testid="stSidebarNavLink"] {{ padding: 3px 10px !important; margin: 1px 0 !important; border: 1px solid transparent; }}
[data-testid="stSidebarNavLink"]:hover {{ background: {rgba(T['accent'], 0.07)} !important; border-color: {T['border']}; }}
[data-testid="stSidebarNavLink"] span {{ font-size: .98rem !important; }}
[data-testid="stSidebarNavLink"][aria-current="page"] {{ border-color: {rgba(T['accent'], 0.35)}; }}
[data-testid="stNavSectionHeader"] {{ font-size: .76rem !important; font-weight: 700 !important; color: {T['accent']} !important;
  margin-top: 14px !important; letter-spacing: .01em; }}
.mx-sideprog {{ border: 1px solid {T['border']}; border-radius: 12px; padding: 10px 12px; margin: 10px 0 12px 0; background: {T['bg']}; }}
.mx-sideprog .top {{ display: flex; justify-content: space-between; font-size: .82rem; color: {T['sub']}; }}
.mx-sideprog .top b {{ color: {T['text']}; }}
.mx-sideprog .mx-bar {{ padding: 2px 0 0 0; }}
.mx-sideprog .nx {{ font-size: .8rem; color: {T['accent']}; margin-top: 7px; font-weight: 600; }}

/* ---------- Bienvenida ---------- */
.mx-hero {{ padding: 6px 0 22px 0; margin-bottom: 20px; border-bottom: 1px solid {T['border']}; position: relative; animation: mdRise .5s ease both; }}
.mx-hero::after {{ content: ""; position: absolute; left: 0; bottom: -1px; height: 2px; width: 160px; background: {T['accent']};
  transform-origin: left; animation: mdDraw .8s .2s cubic-bezier(.2,.7,.2,1) both; box-shadow: {GLOW if NEON else 'none'}; }}
.mx-hero-t {{ font-family: {T['titulos']}; font-size: 2.6rem; font-weight: 700; letter-spacing: -0.025em; line-height: 1.1;
  color: {T['text']}; text-shadow: {GLOW_TXT}; }}
.mx-hero-s {{ color: {T['sub']}; font-size: 1.08rem; margin-top: 8px; max-width: 62ch; line-height: 1.5; }}
.mx-hint {{ display: inline-flex; align-items: center; gap: 8px; margin-top: 16px; padding: 7px 14px; border-radius: 999px;
  border: 1px solid {rgba(T['accent'], 0.3)}; background: {T['soft']}; color: {T['accent']}; font-size: .9rem; font-weight: 600; }}
.mx-hint .ico {{ animation: mdNudge 1.8s ease-in-out infinite; }}
@keyframes mdNudge {{ 0%, 100% {{ transform: translateX(0); }} 50% {{ transform: translateX(-5px); }} }}
.mx-p {{ color: {T['sub']}; line-height: 1.6; margin: 0; font-size: .95rem; }}
.mx-p b {{ color: {T['text']}; }}
.mx-progress .mx-big {{ font-family: {T['titulos']}; font-size: 2.6rem; font-weight: 700; color: {T['accent']}; line-height: 1;
  margin-top: 4px; text-shadow: {GLOW_TXT}; }}
.mx-progress .mx-big span {{ font-size: 1.2rem; margin-left: 2px; }}
.mx-sub {{ color: {T['sub']}; font-size: .88rem; margin-top: 4px; }}
.mx-sub b {{ color: {T['text']}; }}
.mx-empty {{ display: flex; gap: 14px; align-items: flex-start; color: {T['text']}; font-size: .93rem; line-height: 1.5; }}
.mx-empty span {{ color: {T['sub']}; font-size: .86rem; }}

/* Tarjetas de pasos */
.mx-card {{ background: {T['card']}; border: 1px solid {T['border']}; border-radius: 14px; padding: 16px 16px 14px 16px;
  min-height: 205px; margin-bottom: 8px; position: relative; animation: mdRise .45s ease both; animation-delay: var(--d, 0s);
  transition: border-color .2s ease, transform .2s ease; }}
.mx-card:hover {{ border-color: {rgba(T['accent'], 0.6)}; transform: translateY(-2px); }}
.mx-card.next {{ border-color: {T['accent']}; box-shadow: {GLOW}; }}
.mx-card.ok {{ border-color: {rgba(T['ok'], 0.45)}; }}
.mx-card-h {{ display: flex; align-items: center; gap: 10px; }}
.mx-card-h .pill {{ margin-left: auto; }}
.mx-num {{ width: 30px; height: 30px; border-radius: 9px; display: inline-grid; place-items: center; font-weight: 700; font-size: .95rem;
  background: {T['soft']}; color: {T['accent']}; border: 1px solid {rgba(T['accent'], 0.25)}; font-family: {T['titulos']}; }}
.mx-card.next .mx-num {{ background: {T['accent']}; color: {T['on_accent']}; }}
.mx-card.ok .mx-num {{ background: {rgba(T['ok'], 0.14)}; color: {T['ok']}; border-color: {rgba(T['ok'], 0.3)}; }}
.mx-card-i {{ color: {T['sub']}; }}
.mx-card-t {{ font-family: {T['titulos']}; font-weight: 600; font-size: 1.05rem; margin-top: 12px; color: {T['text']}; }}
.mx-card-d {{ color: {T['sub']}; font-size: .88rem; line-height: 1.5; margin-top: 4px; }}
.mx-card-f {{ color: {T['text']}; font-size: .8rem; margin-top: 10px; padding-top: 8px; border-top: 1px dashed {T['border']}; font-weight: 500; }}

.st-key-cta_continuar a {{ background: {T['accent']} !important; border-color: {T['accent']} !important; box-shadow: {GLOW}; }}
.st-key-cta_continuar a span, .st-key-cta_continuar a p {{ color: {T['on_accent']} !important; font-weight: 600; }}
.st-key-cta_continuar a:hover {{ filter: brightness(1.06); }}

/* ---------- Botón colgante (modo oscuro) ---------- */
[data-testid="stAppDeployButton"], [data-testid="stToolbarActions"] {{ display: none !important; }}
.st-key-btn_luna {{ position: fixed !important; top: 0; right: 34px; z-index: 1000001; width: 52px !important;
  display: flex !important; flex-direction: column; align-items: center; transform-origin: top center;
  animation: mdSway 6s ease-in-out infinite; }}
.st-key-btn_luna::before {{ content: ""; display: block; width: 1.5px; height: 30px;
  background: linear-gradient({rgba(T['sub'], 0.2)}, {T['sub']}); }}
.st-key-btn_luna button {{ width: 48px !important; height: 48px !important; min-height: 48px !important; border-radius: 50% !important;
  padding: 0 !important; background: {T['card']} !important; border: 1px solid {T['border']} !important;
  box-shadow: 0 6px 18px {rgba('#000000', 0.12)}, {GLOW if NEON else '0 0 0 0 transparent'} !important; }}
.st-key-btn_luna button:hover {{ border-color: {T['accent']} !important; transform: scale(1.06); }}
.st-key-btn_luna button [data-testid="stIconMaterial"] {{ font-size: 1.45rem !important; color: {T['accent'] if st.session_state.tema == 'claro' else T['text']} !important; }}
.st-key-btn_luna:hover {{ animation-play-state: paused; }}
@keyframes mdSway {{ 0%, 100% {{ transform: rotate(-2.5deg); }} 50% {{ transform: rotate(2.5deg); }} }}

@media (prefers-reduced-motion: reduce) {{
  *, *::before, *::after {{ animation: none !important; transition: none !important; }}
}}
</style>
""", unsafe_allow_html=True)

# =====================================================================
# 5. ICONOS DIBUJADOS (SVG propios, sin emojis)
# =====================================================================
ICONOS = {
    "hub": '<circle cx="12" cy="12" r="3"/><circle cx="5" cy="5" r="2"/><circle cx="19" cy="5" r="2"/><circle cx="5" cy="19" r="2"/>'
           '<circle cx="19" cy="19" r="2"/><path d="m6.5 6.5 3.4 3.4M17.5 6.5l-3.4 3.4M6.5 17.5l3.4-3.4M17.5 17.5l-3.4-3.4"/>',
    "home": '<path d="M4 10.5 12 4l8 6.5V19a1 1 0 0 1-1 1h-4.5v-6h-5v6H5a1 1 0 0 1-1-1v-8.5Z"/>',
    "layers": '<path d="M12 3 3 8l9 5 9-5-9-5Z"/><path d="m3 13 9 5 9-5"/>',
    "nodes": '<rect x="3" y="3" width="7" height="6" rx="1.5"/><rect x="14" y="15" width="7" height="6" rx="1.5"/>'
             '<path d="M6.5 9v3.5a2 2 0 0 0 2 2h5.5"/><circle cx="17.5" cy="6" r="2.5"/><path d="M10 6h5"/>',
    "flow": '<circle cx="6" cy="6" r="2.5"/><circle cx="18" cy="6" r="2.5"/><circle cx="18" cy="18" r="2.5"/>'
            '<path d="M8.5 6h7M6 8.5V13a3 3 0 0 0 3 3h6.5"/>',
    "kanban": '<rect x="3" y="4" width="18" height="16" rx="2"/><path d="M9 4v16M15 4v16"/><path d="M5.5 8h1.5M11 8h2M17 8h1.5M11 11h2"/>',
    "tree": '<rect x="9" y="3" width="6" height="4" rx="1"/><rect x="3" y="17" width="6" height="4" rx="1"/>'
            '<rect x="15" y="17" width="6" height="4" rx="1"/><path d="M12 7v5M6 17v-3a2 2 0 0 1 2-2h8a2 2 0 0 1 2 2v3"/>',
    "curve": '<path d="M3.5 20.5h17"/><path d="M6 17v3.5M10 11v9.5M14 9v11.5M18 15v5.5"/><path d="M3.5 16c3-1 4.5-9 8.5-9.5s5.5 7 8.5 9"/>',
    "shield": '<path d="M12 3 4.5 6v6c0 4.4 3.2 7.6 7.5 9 4.3-1.4 7.5-4.6 7.5-9V6L12 3Z"/><path d="M12 8.5v4.5M12 16h.01"/>',
    "doc": '<path d="M14 3H7a2 2 0 0 0-2 2v14a2 2 0 0 0 2 2h10a2 2 0 0 0 2-2V8l-5-5Z"/><path d="M14 3v5h5M9 13h6M9 17h4"/>',
    "archive": '<rect x="3" y="4" width="18" height="5" rx="1"/><path d="M5 9v10a1 1 0 0 0 1 1h12a1 1 0 0 0 1-1V9M10 13h4"/>',
    "calendar": '<rect x="3.5" y="5" width="17" height="15.5" rx="2"/><path d="M3.5 10h17M8 3v4M16 3v4"/>',
    "clock": '<circle cx="12" cy="12" r="8.5"/><path d="M12 7.5V12l3 2"/>',
    "wallet": '<path d="M5 7V6a2 2 0 0 1 2-2h10v3"/><rect x="3.5" y="7" width="17" height="13" rx="2"/><path d="M16 13.5h2"/>',
    "target": '<circle cx="12" cy="12" r="8.5"/><circle cx="12" cy="12" r="4.5"/><circle cx="12" cy="12" r=".8" fill="currentColor"/>',
    "trend": '<path d="m3 17 6-6 4 4 8-8"/><path d="M15 7h6v6"/>',
    "alert": '<path d="M12 4 2.8 19.5h18.4L12 4Z"/><path d="M12 10v4M12 16.8h.01"/>',
    "check": '<circle cx="12" cy="12" r="8.5"/><path d="m8.5 12.2 2.4 2.4 4.6-5.2"/>',
    "cross": '<circle cx="12" cy="12" r="8.5"/><path d="m9.5 9.5 5 5M14.5 9.5l-5 5"/>',
    "info": '<circle cx="12" cy="12" r="8.5"/><path d="M12 11v5M12 8h.01"/>',
    "users": '<circle cx="9" cy="8" r="3.5"/><path d="M3 20c0-3.3 2.7-6 6-6s6 2.7 6 6"/><path d="M16 4.6a3.5 3.5 0 0 1 0 6.8M18 14.3c1.8.8 3 2.9 3 5.7"/>',
    "list": '<path d="M9 6h11M9 12h11M9 18h11"/><circle cx="4.5" cy="6" r="1"/><circle cx="4.5" cy="12" r="1"/><circle cx="4.5" cy="18" r="1"/>',
    "sliders": '<path d="M4 6h10M18 6h2M4 12h4M12 12h8M4 18h12M20 18h0"/><circle cx="16" cy="6" r="2"/><circle cx="10" cy="12" r="2"/><circle cx="18" cy="18" r="2"/>',
    "grid": '<rect x="3.5" y="3.5" width="7" height="7" rx="1.5"/><rect x="13.5" y="3.5" width="7" height="7" rx="1.5"/>'
            '<rect x="3.5" y="13.5" width="7" height="7" rx="1.5"/><rect x="13.5" y="13.5" width="7" height="7" rx="1.5"/>',
    "flag": '<path d="M5 21V4M5 4h11l-2 4 2 4H5"/>',
    "search": '<circle cx="11" cy="11" r="6.5"/><path d="m16 16 4 4"/>',
    "repeat": '<path d="M4 12a8 8 0 0 1 13.7-5.6L20 9"/><path d="M20 4v5h-5"/><path d="M20 12a8 8 0 0 1-13.7 5.6L4 15"/><path d="M4 20v-5h5"/>',
    "folder": '<path d="M3.5 7.5a2 2 0 0 1 2-2h4l2 2h7a2 2 0 0 1 2 2v8a2 2 0 0 1-2 2h-13a2 2 0 0 1-2-2v-10Z"/>',
    "circle": '<circle cx="12" cy="12" r="7"/>',
    "arrow_left": '<path d="M19 12H5M11 6l-6 6 6 6"/>',
    "gauge": '<path d="M4 17a8 8 0 1 1 16 0"/><path d="m12 17 3.5-5.5"/><path d="M4 17h2M18 17h2M12 9V7"/>',
    "progress": '<circle cx="12" cy="12" r="8.5"/><path d="M12 3.5a8.5 8.5 0 0 1 0 17Z" fill="currentColor" stroke="none" opacity=".35"/>',
}


def ico(nombre, size=18, color="currentColor", stroke=1.7):
    trazo = ICONOS.get(nombre, ICONOS["circle"])
    return (f'<svg class="ico" width="{size}" height="{size}" viewBox="0 0 24 24" fill="none" stroke="{color}" '
            f'stroke-width="{stroke}" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true">{trazo}</svg>')


# =====================================================================
# 6. UTILIDADES DE FORMATO Y COMPONENTES VISUALES
# =====================================================================
def fmt_num(x, dec=0):
    """Formato chileno: punto de miles y coma decimal."""
    try:
        s = f"{float(x):,.{dec}f}"
    except (TypeError, ValueError):
        return str(x)
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def fmt_money(x):
    return "$" + fmt_num(x, 0)


def slug(texto):
    return "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in str(texto).strip()) or "Proyecto"


def html(contenido):
    st.markdown(contenido, unsafe_allow_html=True)


def encabezado(icono, titulo, descripcion=""):
    hoy = date.today().strftime("%d/%m/%Y")
    html(f'<div class="mx-head"><div class="l"><div class="mx-tile">{ico(icono, 22)}</div>'
         f'<div><h1>{titulo}</h1><div class="d">{descripcion}</div></div></div>'
         f'<div class="meta">Proyecto <b>{st.session_state.nombre_proj}</b><br>{hoy}</div></div>')


def seccion(titulo, subtitulo="", icono=None):
    i = ico(icono, 18) if icono else ""
    sub = f'<div class="s">{subtitulo}</div>' if subtitulo else ""
    html(f'<div class="mx-sec">{i}<div><div class="t">{titulo}</div>{sub}</div></div>')


def paso(numero, titulo):
    html(f'<div class="mx-step"><span class="n">{numero}</span>{titulo}</div>')


def franja(celdas):
    """celdas: lista de dicts {icono, etiqueta, valor, sub, destacado}"""
    partes = []
    for i, c in enumerate(celdas):
        clase = "c hl" if c.get("destacado") else "c"
        partes.append(
            f'<div class="{clase}" style="--d:{i * 0.06:.2f}s"><div class="lb">{ico(c.get("icono", "circle"), 15)}'
            f'{c["etiqueta"]}</div><div class="v">{c["valor"]}</div><div class="sb">{c.get("sub", "")}</div></div>')
    html(f'<div class="mx-strip" style="--n:{len(celdas)}">{"".join(partes)}</div>')


def panel(titulo, filas, icono="list", delay=0.0):
    """Panel con filas clave/valor."""
    cuerpo = "".join(f'<div class="mx-row"><span class="k">{k}</span><span class="v">{v}</span></div>' for k, v in filas)
    return f'<div class="mx-panel" style="--d:{delay:.2f}s"><div class="pt">{ico(icono, 17)}{titulo}</div>{cuerpo}</div>'


def barra_prob(etiqueta, valor, meta):
    color = T["ok"] if valor >= meta else (T["warn"] if valor >= 50 else T["bad"])
    return (f'<div class="mx-bar" style="--c:{color};--w:{max(min(valor, 100), 0):.1f}%;--m:{meta}%">'
            f'<div class="top"><span class="k">{etiqueta}</span><span class="v">{fmt_num(valor, 1)}%</span></div>'
            f'<div class="tr"><div class="fl"></div><div class="mk" title="Confianza exigida"></div></div></div>')


def tabla_html(df, resaltar=None):
    cab = "".join(f"<th>{c}</th>" for c in df.columns)
    filas = []
    for fila in df.itertuples(index=False):
        clase = ' class="hl"' if resaltar is not None and resaltar(fila) else ""
        filas.append(f"<tr{clase}>" + "".join(f"<td>{v}</td>" for v in fila) + "</tr>")
    return f'<div class="mx-tw"><table class="mx-table"><thead><tr>{cab}</tr></thead><tbody>{"".join(filas)}</tbody></table></div>'


def pill(texto, tipo, icono=None):
    i = ico(icono, 13, stroke=2) if icono else ""
    return f'<span class="pill {tipo}">{i}{texto}</span>'


def limpiar_widgets(*prefijos):
    """Borra el estado de widgets antiguos."""
    for k in list(st.session_state.keys()):
        if any(str(k).startswith(p) for p in prefijos):
            del st.session_state[k]


def version_lista(nombre):
    """Número de versión de una lista editable. Cambia al eliminar filas para que los
    campos se vuelvan a crear con los valores correctos (evita que la fila siguiente
    muestre el texto de la fila borrada)."""
    return st.session_state.get(f"_ver_{nombre}", 0)


def eliminar_fila(nombre_lista, indice, *prefijos):
    st.session_state[nombre_lista].pop(indice)
    limpiar_widgets(*prefijos)
    st.session_state[f"_ver_{nombre_lista}"] = version_lista(nombre_lista) + 1


def estilo_fig(fig, height=380, **kwargs):
    base = dict(
        paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)",
        font=dict(family=FUENTE, color=T["text"], size=12.5),
        height=height, margin=dict(l=8, r=8, t=36, b=8),
        hoverlabel=dict(bgcolor=T["card"], font_color=T["text"], bordercolor=T["border"], font_family=FUENTE),
        legend=dict(orientation="h", y=1.1, x=0, bgcolor="rgba(0,0,0,0)", font=dict(color=T["sub"])),
        separators=",.",
        transition=dict(duration=400, easing="cubic-in-out"),
    )
    base.update(kwargs)
    fig.update_layout(**base)
    fig.update_xaxes(gridcolor=T["grid"], zeroline=False, linecolor=T["border"], tickfont=dict(color=T["sub"]),
                     title_font=dict(color=T["sub"]))
    fig.update_yaxes(gridcolor=T["grid"], zeroline=False, linecolor=T["border"], tickfont=dict(color=T["sub"]),
                     title_font=dict(color=T["sub"]))
    return fig


PLOTLY_CFG = {"displaylogo": False, "modeBarButtonsToRemove": ["lasso2d", "select2d"]}


def mostrar_fig(fig, key=None):
    st.plotly_chart(fig, config=PLOTLY_CFG, key=key, **ANCHO)


def mostrar_mermaid(codigo, alto="420px", key=None):
    """Renderiza Mermaid con los colores del tema activo (si el código no trae su propio %%{init}%%)."""
    codigo = str(codigo or "").strip()
    if not codigo:
        st.info("Escribe código Mermaid para ver la vista previa.")
        return
    if not codigo.startswith("%%{init"):
        variables = {
            "background": T["card"], "mainBkg": T["card"], "primaryColor": T["card"],
            "primaryTextColor": T["text"], "primaryBorderColor": T["accent"], "lineColor": T["accent"],
            "secondaryColor": T["bg"], "tertiaryColor": T["bg"], "textColor": T["text"],
            "nodeBorder": T["accent"], "clusterBkg": T["bg"], "edgeLabelBackground": T["card"],
            "actorBkg": T["card"], "actorBorder": T["accent"], "actorTextColor": T["text"],
            "actorLineColor": T["sub"], "signalColor": T["text"], "signalTextColor": T["text"],
            "labelBoxBkgColor": T["card"], "labelTextColor": T["text"], "noteBkgColor": T["bg"],
            "noteTextColor": T["text"], "sequenceNumberColor": T["on_accent"],
            "fontFamily": "Consolas, monospace" if NEON else "Segoe UI, Helvetica, Arial, sans-serif",
            "fontSize": "14px",
        }
        init = json.dumps({"theme": "base", "fontFamily": variables["fontFamily"], "themeVariables": variables})
        codigo = f"%%{{init: {init}}}%%\n" + codigo
    with st.container(border=True):
        st_mermaid(codigo, height=alto, key=key)


def generar_mermaid_edt(edt_items, nombre_proyecto):
    """Árbol EDT jerárquico. Los niveles se diferencian por estilo, no por emojis."""
    def limpio(x):
        return str(x).strip().replace('"', '').replace("'", "").replace("$", "")

    lines = ["graph TD", f'    ROOT["<b>{limpio(nombre_proyecto)}</b>"]:::raiz']
    fases = {}
    for item in edt_items:
        fase = limpio(item.get("Fase", "Sin fase")) or "Sin fase"
        subfase = limpio(item.get("Subfase", ""))
        fases.setdefault(fase, {}).setdefault(subfase, []).append({
            "paquete": limpio(item.get("Paquete", "Tarea")),
            "horas": float(item.get("Horas", 0) or 0),
            "costo": float(item.get("Costo", 0) or 0),
            "doc": limpio(item.get("Documentos", "")),
        })

    for fi, (fase_name, subfases) in enumerate(fases.items()):
        fid = f"F{fi}"
        t_h = sum(p["horas"] for s in subfases.values() for p in s)
        t_c = sum(p["costo"] for s in subfases.values() for p in s)
        lines.append(f'    ROOT --> {fid}["<b>{fase_name}</b><br/>{t_h:.0f} h | {fmt_num(t_c)} CLP"]:::fase')
        for si, (sub_name, pkgs) in enumerate(subfases.items()):
            if sub_name:
                sid = f"{fid}_S{si}"
                s_h = sum(p["horas"] for p in pkgs)
                s_c = sum(p["costo"] for p in pkgs)
                lines.append(f'    {fid} --> {sid}["{sub_name}<br/>{s_h:.0f} h | {fmt_num(s_c)} CLP"]:::sub')
                padre = sid
            else:
                padre = fid
            for pi, pkg in enumerate(pkgs):
                det = [f"{pkg['horas']:.0f} h"]
                if pkg["costo"] > 0:
                    det.append(f"{fmt_num(pkg['costo'])} CLP")
                doc = f"<br/><i>{pkg['doc']}</i>" if pkg["doc"] else ""
                lines.append(f'    {padre} --> {padre}_P{pi}["<b>{pkg["paquete"]}</b><br/>{" | ".join(det)}{doc}"]:::pkg')

    lines += [
        f"    classDef raiz fill:{T['accent']},stroke:{T['accent']},color:{T['on_accent']},stroke-width:2px",
        f"    classDef fase fill:{T['soft']},stroke:{T['accent']},color:{T['text']},stroke-width:1.5px",
        f"    classDef sub fill:{T['card']},stroke:{T['sub']},color:{T['text']},stroke-width:1px,stroke-dasharray:4 3",
        f"    classDef pkg fill:{T['card']},stroke:{T['border']},color:{T['text']},stroke-width:1px",
    ]
    return "\n".join(lines)


def clasificar_riesgo(sev):
    if sev <= 5:
        return "Bajo"
    elif sev <= 10:
        return "Medio"
    elif sev <= 15:
        return "Alto"
    return "Crítico"


TIPO_NIVEL = {"Bajo": "ok", "Medio": "warn", "Alto": "bad", "Crítico": "bad"}


def pill_riesgo(sev):
    nivel = clasificar_riesgo(sev)
    return pill(nivel, TIPO_NIVEL[nivel], "alert" if nivel in ("Alto", "Crítico") else None)


def cambiar_plantilla_flujo():
    sel = st.session_state.select_tipo_flujo
    if sel in PLANTILLAS_FLUJO:
        st.session_state.flujo_codigo = PLANTILLAS_FLUJO[sel]


# =====================================================================
# 7. MOTOR MONTE CARLO
# =====================================================================
DISTRIBUCIONES = ["Beta-PERT", "Triangular", "Normal", "Uniforme"]
NIVELES_PERCENTIL = [10, 25, 50, 75, 80, 85, 90, 95, 99]


def _muestra(rng, dist, a, m, b, n):
    """Genera n muestras de una distribución de tres puntos (a=mín, m=más probable, b=máx)."""
    if b <= a:
        return np.full(n, float(m))
    if dist == "Triangular":
        return rng.triangular(a, min(max(m, a), b), b, n)
    if dist == "Uniforme":
        return rng.uniform(a, b, n)
    if dist == "Normal":
        media = (a + 4 * m + b) / 6
        desv = (b - a) / 6
        return np.maximum(rng.normal(media, desv, n), a * 0.5)
    # Beta-PERT (PMBOK)
    alfa = 1 + 4 * (m - a) / (b - a)
    beta = 1 + 4 * (b - m) / (b - a)
    return a + rng.beta(alfa, beta, n) * (b - a)


@st.cache_data(show_spinner=False, max_entries=24)
def simular_manual(dist, p1, p2, p3, n, costo_fijo, costo_diario, var_costo_pct, seed):
    rng = np.random.default_rng(seed)
    if dist == "Normal":  # p1 = media, p2 = desviación
        dias = np.maximum(rng.normal(p1, p2, n), 0.5)
    elif dist == "Uniforme":  # p1 = mín, p3 = máx
        dias = rng.uniform(p1, p3, n)
    else:
        dias = _muestra(rng, dist, p1, p2, p3, n)
    factor = np.maximum(rng.normal(1.0, var_costo_pct / 100.0, n), 0.0)
    costos = costo_fijo + dias * costo_diario * factor
    return {"dias": dias, "costos": costos, "componentes": None, "nombres": None}


@st.cache_data(show_spinner=False, max_entries=24)
def simular_edt(paquetes, horas_dia, pct_opt, pct_pes, dist, n, costo_fijo, seed):
    """Cada paquete de la EDT se simula por separado y se suman (ejecución secuencial)."""
    rng = np.random.default_rng(seed)
    nombres = [p[0] for p in paquetes]
    comp = np.zeros((n, len(paquetes)))
    costos = np.full(n, float(costo_fijo))
    for j, (_, horas, costo) in enumerate(paquetes):
        m = max(float(horas), 0.0) / horas_dia
        if m <= 0:
            continue
        a, b = m * (1 - pct_opt / 100.0), m * (1 + pct_pes / 100.0)
        s = _muestra(rng, dist, a, m, b, n)
        comp[:, j] = s
        costos += float(costo) * (s / m)  # el costo escala con el esfuerzo real
    return {"dias": comp.sum(axis=1), "costos": costos, "componentes": comp, "nombres": nombres}


def calcular_metricas(dias, costos, dias_meta, costo_meta, perc):
    ok_t = dias <= dias_meta
    ok_c = costos <= costo_meta
    m = {
        "n": int(len(dias)),
        "media": float(np.mean(dias)), "desv": float(np.std(dias)),
        "min": float(np.min(dias)), "max": float(np.max(dias)),
        "p50": float(np.percentile(dias, 50)), "pt": float(np.percentile(dias, perc)),
        "p95": float(np.percentile(dias, 95)),
        "p50c": float(np.percentile(costos, 50)), "ptc": float(np.percentile(costos, perc)),
        "p95c": float(np.percentile(costos, 95)), "media_c": float(np.mean(costos)),
        "prob_t": float(ok_t.mean() * 100), "prob_c": float(ok_c.mean() * 100),
        "prob_ambos": float((ok_t & ok_c).mean() * 100),
        "prob_uno": float((ok_t ^ ok_c).mean() * 100),
    }
    m["prob_ninguno"] = 100 - m["prob_ambos"] - m["prob_uno"]
    m["riesgo_t"] = 100 - m["prob_t"]
    m["reserva_dias"] = m["pt"] - m["p50"]
    m["reserva_costo"] = m["ptc"] - m["p50c"]
    m["cv"] = m["desv"] / m["media"] * 100 if m["media"] else 0
    return m


def fecha_fin(inicio, dias, habiles):
    d = int(np.ceil(max(float(dias), 0)))
    if habiles:
        return pd.Timestamp(np.busday_offset(np.datetime64(inicio, "D"), d, roll="forward"))
    return pd.Timestamp(inicio) + pd.Timedelta(days=d)


def tabla_percentiles(dias, costos, inicio, habiles, perc):
    niveles = sorted(set(NIVELES_PERCENTIL + [int(perc)]))
    filas = []
    for p in niveles:
        d = float(np.percentile(dias, p))
        c = float(np.percentile(costos, p))
        filas.append({"Confianza": f"P{p}", "p": p, "Días": d, "Fecha": fecha_fin(inicio, d, habiles), "Costo": c})
    return filas


def veredicto(met, perc):
    pa = met["prob_ambos"]
    if pa >= perc:
        return "ok", "Viable con la confianza exigida"
    if pa >= 50:
        return "warn", "Viable, con riesgo moderado"
    return "bad", "Alto riesgo de incumplimiento"


# =====================================================================
# 8. GENERACIÓN DE DOCUMENTOS (EXCEL / WORD)
# =====================================================================
XL_NAVY, XL_ACC = "1E3A8A", "2563EB"
XL_ZEBRA, XL_SOFT = "F1F5F9", "F8FAFC"
XL_OK, XL_BAD, XL_HL = "DCFCE7", "FEE2E2", "FEF3C7"
XL_BORDE = Border(*(Side(style="thin", color="E2E8F0") for _ in range(4)))


def _xl_titulo(ws, texto, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(ncols, 2))
    c = ws.cell(row=1, column=1, value=texto)
    c.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", start_color=XL_NAVY, end_color=XL_NAVY)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 32


def _xl_meta(ws, fila, pares):
    for etiqueta, valor in pares:
        a = ws.cell(row=fila, column=1, value=etiqueta)
        b = ws.cell(row=fila, column=2, value=valor)
        a.font = Font(name="Calibri", bold=True, color="0F172A")
        a.fill = PatternFill("solid", start_color=XL_SOFT, end_color=XL_SOFT)
        b.font = Font(name="Calibri", color="334155")
        a.border = b.border = XL_BORDE
        fila += 1
    return fila


def _xl_tabla(ws, fila, encabezados, filas, formatos=None, resaltar=None):
    """Escribe una tabla con estilo. formatos: {col_idx(1-based): number_format}. Devuelve la siguiente fila libre."""
    formatos = formatos or {}
    for j, h in enumerate(encabezados, 1):
        c = ws.cell(row=fila, column=j, value=h)
        c.font = Font(name="Calibri", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color=XL_ACC, end_color=XL_ACC)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = XL_BORDE
    ws.row_dimensions[fila].height = 24
    for i, f in enumerate(filas):
        r = fila + 1 + i
        for j, v in enumerate(f, 1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = Font(name="Calibri", color="334155")
            c.border = XL_BORDE
            if j in formatos:
                c.number_format = formatos[j]
            if isinstance(v, (int, float)):
                c.alignment = Alignment(horizontal="right")
            if resaltar is not None and resaltar(i, f):
                c.fill = PatternFill("solid", start_color=XL_HL, end_color=XL_HL)
                c.font = Font(name="Calibri", bold=True, color="0F172A")
            elif i % 2 == 1:
                c.fill = PatternFill("solid", start_color=XL_ZEBRA, end_color=XL_ZEBRA)
    return fila + len(filas) + 1


def _xl_autofit(ws, minimo=12, maximo=70):
    anchos = {}
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.value is None:
                continue
            v = c.value
            if isinstance(v, float) and "%" in (c.number_format or ""):
                txt = f"{v * 100:.1f}%"
            elif isinstance(v, (int, float)) and "$" in (c.number_format or ""):
                txt = f"${v:,.0f}"
            elif isinstance(v, float):
                txt = f"{v:,.2f}"
            else:
                txt = str(v)
            anchos[c.column_letter] = max(anchos.get(c.column_letter, 0), len(txt))
    for col, w in anchos.items():
        ws.column_dimensions[col].width = min(max(w + 4, minimo), maximo)


def _xl_bytes(wb):
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _filas_requisitos():
    filas = [[f"RF-{i:02d}", "Funcional", t.strip()] for i, t in enumerate(st.session_state.rf_list, 1) if t.strip()]
    filas += [[f"RNF-{i:02d}", "No funcional", t.strip()] for i, t in enumerate(st.session_state.rnf_list, 1) if t.strip()]
    return filas


def _filas_riesgos():
    filas = []
    for r in st.session_state.riesgos_list:
        p, i = int(r.get("Probabilidad", 1)), int(r.get("Impacto", 1))
        filas.append([r.get("Riesgo", ""), p, i, p * i, clasificar_riesgo(p * i)])
    return filas


def _filas_edt():
    return [[e.get("Fase", ""), e.get("Subfase", ""), e.get("Paquete", ""), float(e.get("Horas", 0) or 0),
             float(e.get("Costo", 0) or 0), e.get("Documentos", "")] for e in st.session_state.edt_list]


ENC_EDT = ["Fase principal", "Subfase / grupo", "Paquete de trabajo", "Horas estimadas", "Costo estimado ($)", "Documentos"]


def generar_excel_requisitos():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Requisitos"
    _xl_titulo(ws, f"Matriz de requisitos · {st.session_state.nombre_proj}", 3)
    f = _xl_meta(ws, 3, [("Proyecto", st.session_state.nombre_proj), ("Objetivo", st.session_state.objetivo_text)])
    _xl_tabla(ws, f + 1, ["Código", "Tipo de requisito", "Descripción"], _filas_requisitos())
    ws.freeze_panes = ws.cell(row=f + 2, column=1)
    _xl_autofit(ws)
    return _xl_bytes(wb)


def generar_excel_kanban():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tablero Kanban"
    _xl_titulo(ws, f"Tablero Kanban · {st.session_state.nombre_proj}", 3)
    filas = [[i, t["Tarea"], t["Estado"]] for i, t in enumerate(st.session_state.kanban_tasks, 1)]
    _xl_tabla(ws, 3, ["N°", "Tarea", "Estado"], filas)
    ws.freeze_panes = "A4"
    _xl_autofit(ws)
    return _xl_bytes(wb)


def generar_excel_edt():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EDT WBS"
    _xl_titulo(ws, f"Estructura de desglose de trabajo · {st.session_state.nombre_proj}", 6)
    filas = _filas_edt()
    fin = _xl_tabla(ws, 3, ENC_EDT, filas, formatos={4: "#,##0.0", 5: "$#,##0"})
    if filas:
        ws.cell(row=fin, column=3, value="TOTAL").font = Font(bold=True)
        ws.cell(row=fin, column=4, value=f"=SUM(D4:D{fin - 1})").number_format = "#,##0.0"
        ws.cell(row=fin, column=5, value=f"=SUM(E4:E{fin - 1})").number_format = "$#,##0"
        ws.cell(row=fin, column=4).font = ws.cell(row=fin, column=5).font = Font(bold=True)
    ws.freeze_panes = "A4"
    _xl_autofit(ws)
    return _xl_bytes(wb)


def generar_excel_riesgos():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Matriz de Riesgos"
    _xl_titulo(ws, f"Matriz de riesgos · {st.session_state.nombre_proj}", 5)
    _xl_tabla(ws, 3, ["Riesgo", "Probabilidad (1-5)", "Impacto (1-5)", "Severidad (P×I)", "Nivel"],
              _filas_riesgos(), resaltar=lambda i, f: f[3] >= 16)
    ws.freeze_panes = "A4"
    _xl_autofit(ws)
    return _xl_bytes(wb)


@st.cache_data(show_spinner=False, max_entries=8)
def generar_excel_montecarlo(dias, costos, cfg):
    wb = openpyxl.Workbook()
    met = calcular_metricas(dias, costos, cfg["dias_meta"], cfg["costo_meta"], cfg["perc"])
    perc = cfg["perc"]

    # --- Hoja 1: Resumen ejecutivo
    ws = wb.active
    ws.title = "Resumen"
    _xl_titulo(ws, f"Simulación Monte Carlo · {cfg['proyecto']}", 3)
    f = _xl_meta(ws, 3, [
        ("Proyecto", cfg["proyecto"]),
        ("Fuente de estimación", cfg["fuente"]),
        ("Distribución", cfg["dist"]),
        ("Parámetros", cfg["parametros"]),
        ("Iteraciones", met["n"]),
        ("Fecha de inicio", cfg["inicio"]),
        ("Calendario", "Días hábiles" if cfg["habiles"] else "Días corridos"),
        ("Plazo meta (días)", cfg["dias_meta"]),
        ("Presupuesto meta ($)", cfg["costo_meta"]),
        ("Confianza exigida", f"P{perc}"),
    ])
    ws.cell(row=f - 2, column=2).number_format = "#,##0.0"
    ws.cell(row=f - 1, column=2).number_format = "$#,##0"
    _, titulo_v = veredicto(met, perc)
    f += 1
    ws.cell(row=f, column=1, value="Veredicto").font = Font(bold=True, size=12, color="0F172A")
    ws.cell(row=f, column=2, value=titulo_v).font = Font(bold=True, size=12, color=XL_NAVY)
    f += 2

    filas_t = [
        ["Duración mediana (P50)", met["p50"], "días"],
        [f"Duración comprometida (P{perc})", met["pt"], "días"],
        ["Duración pesimista (P95)", met["p95"], "días"],
        ["Duración promedio", met["media"], "días"],
        ["Desviación estándar", met["desv"], "días"],
        ["Reserva de contingencia de tiempo", met["reserva_dias"], "días"],
        ["Probabilidad de cumplir el plazo", met["prob_t"] / 100, "%"],
        ["Probabilidad de cumplir el presupuesto", met["prob_c"] / 100, "%"],
        ["Probabilidad de cumplir ambos", met["prob_ambos"] / 100, "%"],
    ]
    fmt_t = {}
    f_ini = f
    f = _xl_tabla(ws, f, ["Indicador de tiempo y éxito", "Valor", "Unidad"], filas_t,
                  resaltar=lambda i, r: "comprometida" in r[0] or "ambos" in r[0])
    for i, r in enumerate(filas_t):
        ws.cell(row=f_ini + 1 + i, column=2).number_format = "0.0%" if r[2] == "%" else "#,##0.0"
        if r[2] == "%":
            v = r[1]
            ws.cell(row=f_ini + 1 + i, column=2).fill = PatternFill(
                "solid", start_color=XL_OK if v >= perc / 100 else (XL_BAD if v < 0.5 else XL_HL))
    f += 1
    filas_c = [
        ["Costo mediano (P50)", met["p50c"]],
        [f"Costo recomendado (P{perc})", met["ptc"]],
        ["Costo pesimista (P95)", met["p95c"]],
        ["Reserva de contingencia de costo", met["reserva_costo"]],
    ]
    _xl_tabla(ws, f, ["Indicador financiero", "Valor ($)"], filas_c, formatos={2: "$#,##0"},
              resaltar=lambda i, r: "recomendado" in r[0])
    _xl_autofit(ws)
    ws.column_dimensions["A"].width = 42

    # --- Hoja 2: Percentiles
    ws2 = wb.create_sheet("Percentiles")
    _xl_titulo(ws2, "Tabla de confianza (percentiles)", 4)
    tp = tabla_percentiles(dias, costos, pd.Timestamp(cfg["inicio"]), cfg["habiles"], perc)
    _xl_tabla(ws2, 3, ["Confianza", "Días", "Fecha estimada de término", "Costo ($)"],
              [[r["Confianza"], r["Días"], r["Fecha"].strftime("%d/%m/%Y"), r["Costo"]] for r in tp],
              formatos={2: "#,##0.0", 4: "$#,##0"}, resaltar=lambda i, r: r[0] == f"P{perc}")
    _xl_autofit(ws2)

    # --- Hoja 3: Histograma con gráfico nativo de Excel
    ws3 = wb.create_sheet("Histograma")
    _xl_titulo(ws3, "Distribución de duraciones simuladas", 4)
    conteos, bordes = np.histogram(dias, bins=30)
    filas_h = [[f"{bordes[i]:.1f} – {bordes[i + 1]:.1f}", int(conteos[i]), float(conteos[i] / len(dias)),
                float((dias <= bordes[i + 1]).mean())] for i in range(len(conteos))]
    _xl_tabla(ws3, 3, ["Rango (días)", "Escenarios", "% del total", "% acumulado"], filas_h,
              formatos={3: "0.0%", 4: "0.0%"})
    _xl_autofit(ws3)
    graf = BarChart()
    graf.title = "Frecuencia de duraciones"
    graf.y_axis.title = "Escenarios"
    graf.x_axis.title = "Días"
    graf.add_data(Reference(ws3, min_col=2, min_row=3, max_row=3 + len(filas_h)), titles_from_data=True)
    graf.set_categories(Reference(ws3, min_col=1, min_row=4, max_row=3 + len(filas_h)))
    graf.gapWidth = 10
    graf.height, graf.width = 9, 18
    graf.legend = None
    ws3.add_chart(graf, "F3")
    curva = LineChart()
    curva.title = "Curva S (probabilidad acumulada)"
    curva.add_data(Reference(ws3, min_col=4, min_row=3, max_row=3 + len(filas_h)), titles_from_data=True)
    curva.set_categories(Reference(ws3, min_col=1, min_row=4, max_row=3 + len(filas_h)))
    curva.height, curva.width = 9, 18
    curva.legend = None
    ws3.add_chart(curva, "F23")

    # --- Hoja 4: Datos crudos
    ws4 = wb.create_sheet("Simulaciones")
    ws4.append(["Iteración", "Duración (días)", "Costo ($)", "Cumple plazo", "Cumple presupuesto"])
    for c in ws4[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color=XL_ACC, end_color=XL_ACC)
    for i, (d, c) in enumerate(zip(dias, costos), 1):
        ws4.append([i, round(float(d), 2), round(float(c), 0),
                    "Sí" if d <= cfg["dias_meta"] else "No", "Sí" if c <= cfg["costo_meta"] else "No"])
    for col, w in zip("ABCDE", [12, 18, 16, 14, 20]):
        ws4.column_dimensions[col].width = w
    ws4.freeze_panes = "A2"
    return _xl_bytes(wb)


def generar_excel_estilizado():
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Requisitos"
    _xl_titulo(ws1, f"Requisitos · {st.session_state.nombre_proj}", 3)
    _xl_tabla(ws1, 3, ["Código", "Tipo", "Descripción"], _filas_requisitos())
    _xl_autofit(ws1)

    ws2 = wb.create_sheet("Kanban")
    _xl_titulo(ws2, "Tablero Kanban", 2)
    _xl_tabla(ws2, 3, ["Tarea", "Estado"], [[t["Tarea"], t["Estado"]] for t in st.session_state.kanban_tasks])
    _xl_autofit(ws2)

    ws3 = wb.create_sheet("EDT (WBS)")
    _xl_titulo(ws3, "Estructura de desglose de trabajo", 6)
    _xl_tabla(ws3, 3, ENC_EDT, _filas_edt(), formatos={4: "#,##0.0", 5: "$#,##0"})
    _xl_autofit(ws3)

    ws4 = wb.create_sheet("Riesgos")
    _xl_titulo(ws4, "Matriz de riesgos", 5)
    _xl_tabla(ws4, 3, ["Riesgo", "Probabilidad", "Impacto", "Severidad", "Nivel"], _filas_riesgos(),
              resaltar=lambda i, f: f[3] >= 16)
    _xl_autofit(ws4)

    if MC_ACTUAL:
        met, cfg = MC_ACTUAL["met"], MC_ACTUAL["cfg"]
        ws5 = wb.create_sheet("Monte Carlo")
        _xl_titulo(ws5, "Resumen Monte Carlo", 2)
        _xl_tabla(ws5, 3, ["Indicador", "Valor"], [
            ["Distribución", cfg["dist"]], ["Fuente", cfg["fuente"]], ["Iteraciones", met["n"]],
            ["P50 (días)", round(met["p50"], 1)], [f"P{cfg['perc']} (días)", round(met["pt"], 1)],
            ["P95 (días)", round(met["p95"], 1)], [f"Costo P{cfg['perc']} ($)", round(met["ptc"])],
            ["Prob. cumplir plazo", f"{met['prob_t']:.1f}%"], ["Prob. cumplir presupuesto", f"{met['prob_c']:.1f}%"],
            ["Prob. cumplir ambos", f"{met['prob_ambos']:.1f}%"],
        ])
        _xl_autofit(ws5)
    return _xl_bytes(wb)


def tabla_word(doc, encabezados, filas):
    t = doc.add_table(rows=1, cols=len(encabezados))
    try:
        t.style = "Light Grid Accent 1"
    except Exception:
        t.style = "Table Grid"
    for i, h in enumerate(encabezados):
        celda = t.rows[0].cells[i]
        celda.text = ""
        celda.paragraphs[0].add_run(str(h)).bold = True
    for f in filas:
        cells = t.add_row().cells
        for i, v in enumerate(f):
            cells[i].text = str(v)
    doc.add_paragraph()
    return t


def parrafo_analisis(doc, titulo, texto):
    p = doc.add_paragraph()
    r = p.add_run(titulo + " ")
    r.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    p.add_run(texto)


def _doc_base():
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)
    return doc


def generar_word_ers():
    doc = _doc_base()
    doc.add_heading(f"Especificación de requisitos: {st.session_state.nombre_proj}", level=0)
    doc.add_paragraph(f"Equipo: {st.session_state.integrantes}")
    doc.add_paragraph(f"Objetivo: {st.session_state.objetivo_text}")
    doc.add_heading("Matriz de requisitos", level=1)
    filas = _filas_requisitos()
    if filas:
        tabla_word(doc, ["Código", "Tipo", "Descripción"], filas)
    else:
        doc.add_paragraph("No hay requisitos registrados.")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def generar_word_mementomori():
    doc = _doc_base()
    nombre = st.session_state.nombre_proj

    doc.add_heading(f"Reporte Memento Mori: {nombre}", level=0)
    doc.add_paragraph(f"Integrantes / responsables: {st.session_state.integrantes}")
    doc.add_paragraph("Análisis crítico integral de requisitos, arquitectura, tiempos, EDT y riesgos del sistema.")

    # 1. Introducción
    doc.add_heading("1. Introducción y objetivos del proyecto", level=1)
    doc.add_paragraph(f"Objetivo principal: {st.session_state.objetivo_text}")
    parrafo_analisis(doc, "Análisis crítico e interpretación:",
        f"El objetivo planteado para el proyecto '{nombre}' aborda la necesidad clave de automatización e integración. "
        "Sin embargo, el punto crítico principal radica en controlar el alcance funcional inicial. "
        "Bajo el análisis Memento Mori, el principal riesgo de esta fase es la expansión descontrolada del alcance (scope creep), "
        "la cual puede desalinear las expectativas de los interesados antes de consolidar el producto mínimo viable (MVP).")

    # 2. Requisitos
    doc.add_heading("2. Especificación de requisitos (RF y RNF)", level=1)
    filas = _filas_requisitos()
    if filas:
        tabla_word(doc, ["Código", "Tipo", "Descripción"], filas)
    parrafo_analisis(doc, "Análisis crítico de requisitos:",
        "Los puntos críticos se concentran en los requisitos no funcionales (RNF). "
        "Si bien los requisitos funcionales aseguran la operatividad del software, un fallo en el cumplimiento de los RNF "
        "(latencias elevadas o vulnerabilidades de seguridad) puede provocar el rechazo de los usuarios finales "
        "e impedir la escalabilidad del sistema.")

    # 3. Arquitectura
    doc.add_heading("3. Arquitectura y diagramas de flujo / secuencia", level=1)
    doc.add_paragraph("Sintaxis del diagrama activo:")
    p_code = doc.add_paragraph()
    r_code = p_code.add_run(st.session_state.flujo_codigo)
    r_code.font.name = "Consolas"
    r_code.font.size = Pt(9)
    parrafo_analisis(doc, "Interpretación de arquitectura y puntos críticos:",
        "Los diagramas permiten entender la interacción de componentes y el recorrido de la información. "
        "El punto crítico es el acoplamiento directo entre el controlador y los servicios de base de datos o autenticación. "
        "Sin políticas de reintento, caché o manejo de excepciones HTTP (401, 500), una caída puntual de la base de datos "
        "provocará un colapso en cadena en los endpoints expuestos al usuario.")

    # 4. Kanban
    doc.add_heading("4. Estado de gestión de tareas (tablero Kanban)", level=1)
    tareas = st.session_state.kanban_tasks
    if tareas:
        tabla_word(doc, ["Tarea", "Estado"], [[t["Tarea"], t["Estado"]] for t in tareas])
        n_proc = sum(1 for t in tareas if t["Estado"] == "En Proceso")
        n_ok = sum(1 for t in tareas if t["Estado"] == "Completado")
        doc.add_paragraph(f"Avance: {n_ok} de {len(tareas)} tareas completadas ({n_ok / len(tareas) * 100:.0f}%). "
                          f"Trabajo en proceso (WIP): {n_proc} tareas.")
    else:
        doc.add_paragraph("No existen tareas registradas.")
    parrafo_analisis(doc, "Análisis crítico del flujo Kanban:",
        "La acumulación de ítems 'En Proceso' es la mayor amenaza para el ritmo de entrega (WIP). "
        "Se sugiere definir un límite estricto de tareas simultáneas para evitar sobrecarga y sobreestimación del rendimiento.")

    # 5. EDT
    doc.add_heading("5. Estructura de desglose de trabajo (EDT / WBS)", level=1)
    edt = st.session_state.edt_list
    if edt:
        tabla_word(doc, ["Fase", "Paquete de trabajo", "Horas", "Costo ($)", "Documentos"],
                   [[e.get("Fase", ""), e.get("Paquete", ""), fmt_num(e.get("Horas", 0), 1),
                     fmt_money(e.get("Costo", 0)), e.get("Documentos", "")] for e in edt])
        tot_h = sum(float(e.get("Horas", 0) or 0) for e in edt)
        tot_c = sum(float(e.get("Costo", 0) or 0) for e in edt)
        doc.add_paragraph(f"Totales planificados: {fmt_num(tot_h, 1)} horas | Presupuesto: {fmt_money(tot_c)}")
    else:
        doc.add_paragraph("No hay paquetes de EDT registrados.")
    parrafo_analisis(doc, "Puntos críticos presupuestarios y de carga horaria:",
        "Las actividades con mayor asignación de costo y horas conforman la ruta crítica. Cualquier imprevisto de estimación "
        "en ellas provoca desviaciones acumulativas que agotan la reserva de contingencia y extienden la fecha límite.")

    # 6. Monte Carlo
    doc.add_heading("6. Análisis y simulación Monte Carlo", level=1)
    if MC_ACTUAL:
        met, cfg = MC_ACTUAL["met"], MC_ACTUAL["cfg"]
        perc = cfg["perc"]
        doc.add_paragraph(f"Fuente de estimación: {cfg['fuente']} | Distribución: {cfg['dist']} | "
                          f"Iteraciones: {fmt_num(met['n'])} | Parámetros: {cfg['parametros']}")
        doc.add_paragraph(f"Metas evaluadas: plazo de {fmt_num(cfg['dias_meta'], 1)} días y presupuesto de "
                          f"{fmt_money(cfg['costo_meta'])}, con confianza exigida P{perc}.")
        tabla_word(doc, ["Indicador", "Valor"], [
            ["Duración mediana (P50)", f"{fmt_num(met['p50'], 1)} días"],
            [f"Duración comprometida (P{perc})", f"{fmt_num(met['pt'], 1)} días"],
            ["Duración pesimista (P95)", f"{fmt_num(met['p95'], 1)} días"],
            ["Reserva de contingencia (tiempo)", f"{fmt_num(met['reserva_dias'], 1)} días"],
            [f"Costo recomendado (P{perc})", fmt_money(met["ptc"])],
            ["Reserva de contingencia (costo)", fmt_money(met["reserva_costo"])],
            ["Probabilidad de cumplir el plazo", f"{fmt_num(met['prob_t'], 1)}%"],
            ["Probabilidad de cumplir el presupuesto", f"{fmt_num(met['prob_c'], 1)}%"],
            ["Probabilidad de cumplir ambos", f"{fmt_num(met['prob_ambos'], 1)}%"],
        ])
        _, titulo_v = veredicto(met, perc)
        parrafo_analisis(doc, f"Veredicto: {titulo_v}.",
            f"De cada 100 escenarios simulados, {round(met['prob_ambos'])} terminan dentro de plazo y presupuesto. "
            f"Para comprometer la entrega con un {perc}% de confianza se requieren {fmt_num(met['pt'], 1)} días "
            f"(fecha estimada {MC_ACTUAL['fecha_pt'].strftime('%d/%m/%Y')}) y un presupuesto de {fmt_money(met['ptc'])}. "
            f"La diferencia frente a la mediana ({fmt_num(met['reserva_dias'], 1)} días y {fmt_money(met['reserva_costo'])}) "
            "corresponde a la reserva de contingencia recomendada. El escenario pesimista (P95) "
            f"alcanza {fmt_num(met['p95'], 1)} días, por lo que se aconsejan controles semanales de seguimiento.")
    else:
        doc.add_paragraph("Los parámetros configurados para la simulación Monte Carlo no son válidos.")

    # 7. Riesgos
    doc.add_heading("7. Evaluación y filosofía de riesgos (Memento Mori)", level=1)
    filas_r = _filas_riesgos()
    if filas_r:
        tabla_word(doc, ["Riesgo", "Probabilidad", "Impacto", "Severidad", "Nivel"], filas_r)
    else:
        doc.add_paragraph("No hay riesgos registrados en la matriz.")
    parrafo_analisis(doc, "Estrategia de mitigación Memento Mori:",
        "Bajo el principio 'Memento Mori' (recordar la vulnerabilidad y planificar anticipadamente los fallos), "
        "los riesgos de severidad alta y crítica deben contar con planes de contingencia documentados antes de iniciar el desarrollo.")

    # 8. Conclusiones
    doc.add_heading("8. Conclusiones y recomendaciones finales", level=1)
    doc.add_paragraph(
        f"El proyecto '{nombre}' cuenta con una base metodológica y técnica sólida. Se recomienda formalizar el compromiso "
        "de entrega según la simulación Monte Carlo y monitorear de forma continua los requisitos no funcionales y la matriz "
        "de riesgos para asegurar el éxito del proyecto.")

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# =====================================================================
# 9. EJECUCIÓN DE LA SIMULACIÓN (se calcula siempre, aunque no se esté en la página)
# =====================================================================
def ejecutar_mc():
    """Lee los parámetros guardados y ejecuta la simulación. Devuelve (resultado | None, mensaje_error)."""
    ss = st.session_state
    desde_edt = str(ss.mc_fuente).startswith("Desde")
    dist = ss.mc_dist_type
    n = int(ss.mc_sims)
    seed = 42 if ss.mc_seed else int(ss.mc_semilla_val)

    if desde_edt:
        paquetes = tuple((str(e.get("Paquete", "")) or f"Paquete {i + 1}", float(e.get("Horas", 0) or 0),
                          float(e.get("Costo", 0) or 0)) for i, e in enumerate(ss.edt_list))
        if not paquetes or sum(p[1] for p in paquetes) <= 0:
            return None, "La EDT no tiene paquetes con horas. Agrégalos en el módulo EDT o usa la estimación manual."
        res = simular_edt(paquetes, float(ss.mc_horas_dia), float(ss.mc_pct_opt), float(ss.mc_pct_pes), dist,
                          n, float(ss.mc_cfijo), seed)
        parametros = f"{fmt_num(ss.mc_horas_dia, 1)} h/día, mejor caso −{ss.mc_pct_opt}%, peor caso +{ss.mc_pct_pes}%"
    else:
        cf, cd, vc = float(ss.mc_cfijo), float(ss.mc_cdiario), float(ss.mc_var_costo)
        if dist in ("Triangular", "Beta-PERT"):
            o, m, p = float(ss.mc_opt), float(ss.mc_prob), float(ss.mc_pes)
            if not (o <= m <= p) or o >= p:
                return None, "Revisa los valores: optimista ≤ más probable ≤ pesimista, y optimista < pesimista."
            res = simular_manual(dist, o, m, p, n, cf, cd, vc, seed)
            parametros = f"O={fmt_num(o, 1)}, M={fmt_num(m, 1)}, P={fmt_num(p, 1)} días"
        elif dist == "Normal":
            res = simular_manual(dist, float(ss.mc_mean), float(ss.mc_std), 0.0, n, cf, cd, vc, seed)
            parametros = f"media={fmt_num(ss.mc_mean, 1)}, desviación={fmt_num(ss.mc_std, 1)} días"
        else:
            a, b = float(ss.mc_min_u), float(ss.mc_max_u)
            if a >= b:
                return None, "En la distribución uniforme el mínimo debe ser menor que el máximo."
            res = simular_manual(dist, a, 0.0, b, n, cf, cd, vc, seed)
            parametros = f"mín={fmt_num(a, 1)}, máx={fmt_num(b, 1)} días"

    perc = int(ss.mc_target_perc)
    dias_meta, costo_meta = float(ss.mc_dias_meta), float(ss.mc_costo_meta)
    inicio, habiles = ss.mc_f_inicio, bool(ss.mc_habiles)
    met = calcular_metricas(res["dias"], res["costos"], dias_meta, costo_meta, perc)
    cfg = {
        "proyecto": ss.nombre_proj, "fuente": ss.mc_fuente, "dist": dist, "parametros": parametros,
        "dias_meta": dias_meta, "costo_meta": costo_meta, "perc": perc, "inicio": str(inicio), "habiles": habiles,
    }
    return {
        "res": res, "met": met, "cfg": cfg,
        "fecha_p50": fecha_fin(inicio, met["p50"], habiles),
        "fecha_pt": fecha_fin(inicio, met["pt"], habiles),
        "fecha_p95": fecha_fin(inicio, met["p95"], habiles),
    }, None


MC_ACTUAL, MC_ERROR = ejecutar_mc()


def nuevas_semilla():
    st.session_state.mc_semilla_val = int(np.random.default_rng().integers(1, 10**9))


# =====================================================================
# 10. PÁGINAS
# =====================================================================
# ---------------------------------------------------------------------
# Guía paso a paso y avance del proyecto
# ---------------------------------------------------------------------
PASOS = [
    ("requisitos", "Requisitos", "layers",
     "Escribe el objetivo del proyecto y define qué debe hacer el sistema (requisitos funcionales) "
     "y cómo debe comportarse (no funcionales).", True),
    ("modelado", "Modelado", "nodes",
     "Dibuja la arquitectura con un diagrama Mermaid. Parte del ejemplo y edítalo a tu medida.", False),
    ("flujos", "Flujos", "flow",
     "Elige una plantilla de flujo o de secuencia (login, API) y adáptala a tu sistema.", False),
    ("kanban", "Kanban", "kanban",
     "Registra las tareas del equipo y su estado: pendiente, en proceso o completada.", True),
    ("edt", "EDT", "tree",
     "Desglosa el trabajo en fases y paquetes, cada uno con horas y costo estimados.", True),
    ("montecarlo", "Monte Carlo", "curve",
     "Simula miles de escenarios y descubre qué plazo y presupuesto puedes comprometer.", True),
    ("riesgos", "Riesgos", "shield",
     "Anota los riesgos con su probabilidad e impacto y revísalos en el mapa de calor.", True),
    ("reporte", "Reporte Memento Mori", "doc",
     "Descarga el informe crítico en Word, armado con todo lo que ingresaste.", False),
    ("descargas", "Descarga completa", "archive",
     "Baja todos los entregables (Word, Excel, CSV y diagramas) en un solo ZIP.", False),
]


def marcar_visita(clave):
    st.session_state.setdefault("_visto", set()).add(clave)


def marcar_descarga(clave):
    st.session_state.setdefault("_descargado", set()).add(clave)


def estado_paso(clave):
    """Devuelve (completo, detalle) de cada paso según los datos ingresados."""
    ss = st.session_state
    visto = ss.get("_visto", set())
    bajado = ss.get("_descargado", set())
    if clave == "requisitos":
        ok = bool(str(ss.objetivo_text).strip()) and bool(ss.rf_list) and bool(ss.rnf_list)
        return ok, f"{len(ss.rf_list)} funcionales y {len(ss.rnf_list)} no funcionales"
    if clave == "modelado":
        return "modelado" in visto, "Diagrama revisado" if "modelado" in visto else "Sin revisar"
    if clave == "flujos":
        return "flujos" in visto, "Flujo revisado" if "flujos" in visto else "Sin revisar"
    if clave == "kanban":
        n = len(ss.kanban_tasks)
        return n > 0, f"{n} tarea(s) registradas"
    if clave == "edt":
        h = sum(float(e.get("Horas", 0) or 0) for e in ss.edt_list)
        return len(ss.edt_list) > 0 and h > 0, f"{len(ss.edt_list)} paquete(s), {fmt_num(h)} h"
    if clave == "montecarlo":
        ok = "montecarlo" in visto and MC_ACTUAL is not None
        return ok, (f"Plazo P{MC_ACTUAL['cfg']['perc']}: {fmt_num(MC_ACTUAL['met']['pt'], 1)} días" if ok else "Sin simular")
    if clave == "riesgos":
        n = len(ss.riesgos_list)
        return n > 0, f"{n} riesgo(s) registrados"
    if clave == "reporte":
        return "reporte" in bajado, "Descargado" if "reporte" in bajado else "Sin descargar"
    if clave == "descargas":
        return "descargas" in bajado, "Descargado" if "descargas" in bajado else "Sin descargar"
    return False, ""


def avance_proyecto():
    estados = [(i + 1, p, *estado_paso(p[0])) for i, p in enumerate(PASOS)]
    hechos = sum(1 for e in estados if e[2])
    siguiente = next((e for e in estados if not e[2]), None)
    return estados, hechos, siguiente


def cargar_ejemplo():
    import copy
    ss = st.session_state
    for k, v in DATOS_EJEMPLO.items():
        ss[k] = copy.deepcopy(v)
    for lista in ("rf_list", "rnf_list", "kanban_tasks", "edt_list", "riesgos_list"):
        ss[f"_ver_{lista}"] = ss.get(f"_ver_{lista}", 0) + 1


def vaciar_proyecto():
    ss = st.session_state
    for lista in ("rf_list", "rnf_list", "kanban_tasks", "edt_list", "riesgos_list"):
        ss[lista] = []
        ss[f"_ver_{lista}"] = ss.get(f"_ver_{lista}", 0) + 1
    ss.objetivo_text = ""
    ss.integrantes = ""
    ss.nombre_proj = "MatrixDev Core"
    ss._visto = set()
    ss._descargado = set()


def pie_navegacion(clave):
    """Enlaces al paso anterior y al siguiente, al final de cada página."""
    claves = [p[0] for p in PASOS]
    if clave not in claves:
        return
    i = claves.index(clave)
    st.write("")
    st.divider()
    c1, c2, c3 = st.columns([1, 1, 1])
    if i > 0:
        c1.page_link(PAGINAS[claves[i - 1]], label=f"Anterior: {i}. {PASOS[i - 1][1]}", icon=":material/arrow_back:", **ANCHO)
    else:
        c1.page_link(PG_INICIO, label="Volver al inicio", icon=":material/arrow_back:", **ANCHO)
    c2.page_link(PG_INICIO, label="Ver avance del proyecto", icon=":material/space_dashboard:", **ANCHO)
    if i < len(claves) - 1:
        c3.page_link(PAGINAS[claves[i + 1]], label=f"Siguiente: {i + 2}. {PASOS[i + 1][1]}", icon=":material/arrow_forward:", **ANCHO)


def pagina_inicio():
    ss = st.session_state
    estados, hechos, siguiente = avance_proyecto()
    pct = hechos / len(PASOS) * 100

    # ---------------- Presentación ----------------
    html(f'<div class="mx-hero"><div class="mx-hero-t">{ss.nombre_proj}</div>'
         f'<div class="mx-hero-s">Tu espacio para planificar, estimar y documentar un proyecto de software, '
         f'desde los requisitos hasta el informe final.</div>'
         f'<div class="mx-hint">{ico("arrow_left", 17)}Empieza en el menú de la izquierda y avanza punto por punto</div></div>')

    c_info, c_prog = st.columns([1.5, 1], gap="large")
    with c_info:
        html(f'<div class="mx-panel"><div class="pt">{ico("info", 17)}¿Qué es {ss.nombre_proj}?</div>'
             f'<p class="mx-p">Es una herramienta que reúne en un solo lugar las piezas de la gestión de un proyecto de software: '
             f'requisitos, diagramas, tareas, desglose del trabajo, estimación probabilística de plazos y costos, y riesgos.</p>'
             f'<div class="pt" style="margin-top:14px">{ico("target", 17)}¿Para qué sirve?</div>'
             f'<p class="mx-p">Para responder con datos tres preguntas clave: <b>qué</b> hay que construir, <b>cuánto</b> '
             f'tiempo y dinero costará con un nivel de confianza conocido, y <b>qué</b> podría salir mal. '
             f'Al final genera el informe en Word y las planillas en Excel listas para entregar.</p></div>')
    with c_prog:
        sig_txt = (f"{siguiente[0]}. {siguiente[1][1]}" if siguiente else "Todo listo")
        html(f'<div class="mx-panel mx-progress"><div class="pt">{ico("progress", 17)}Avance del proyecto</div>'
             f'<div class="mx-big">{fmt_num(pct)}<span>%</span></div>'
             f'<div class="mx-sub">{hechos} de {len(PASOS)} pasos completados</div>'
             f'<div class="mx-bar" style="--c:{T["accent"]};--w:{pct:.1f}%;--m:101%"><div class="tr"><div class="fl"></div></div></div>'
             f'<div class="mx-sub" style="margin-top:10px">Siguiente paso: <b>{sig_txt}</b></div></div>')
        if siguiente:
            with st.container(key="cta_continuar"):
                st.page_link(PAGINAS[siguiente[1][0]], label=f"Continuar con el paso {siguiente[0]}: {siguiente[1][1]}",
                             icon=":material/arrow_forward:", **ANCHO)
        b1, b2 = st.columns(2)
        b1.button("Cargar ejemplo", icon=":material/auto_awesome:", on_click=cargar_ejemplo,
                  help="Rellena el proyecto con datos de muestra para ver cómo funciona todo.", **ANCHO)
        b2.button("Empezar de cero", icon=":material/restart_alt:", on_click=vaciar_proyecto,
                  help="Borra todos los datos ingresados.", **ANCHO)

    # ---------------- Cómo empezar ----------------
    st.write("")
    seccion("Cómo empezar", "Abre el menú de la izquierda y recorre los pasos en orden. Cada paso alimenta al siguiente "
                            "y este inicio se va completando solo con lo que ingresas.", "flag")
    for fila in range(0, len(estados), 3):
        cols = st.columns(3, gap="medium")
        for col, (num, (clave, titulo, icono, desc, _req), ok, detalle) in zip(cols, estados[fila:fila + 3]):
            es_sig = siguiente is not None and siguiente[0] == num
            clase = "mx-card ok" if ok else ("mx-card next" if es_sig else "mx-card")
            estado = pill("Completo", "ok", "check") if ok else (pill("Siguiente", "warn") if es_sig else pill("Pendiente", "neu"))
            with col:
                html(f'<div class="{clase}" style="--d:{(num - 1) * 0.04:.2f}s"><div class="mx-card-h">'
                     f'<span class="mx-num">{num}</span><span class="mx-card-i">{ico(icono, 18)}</span>{estado}</div>'
                     f'<div class="mx-card-t">{titulo}</div><div class="mx-card-d">{desc}</div>'
                     f'<div class="mx-card-f">{detalle}</div></div>')
                st.page_link(PAGINAS[clave], label=f"Abrir paso {num}", icon=":material/arrow_forward:", **ANCHO)

    # ---------------- Resumen que se va calculando ----------------
    st.write("")
    seccion("Resumen del proyecto", "Se calcula con los datos que vas ingresando en cada paso.", "grid")
    rf, rnf = ss.rf_list, ss.rnf_list
    tareas, edt, riesgos = ss.kanban_tasks, ss.edt_list, ss.riesgos_list
    n_ok = sum(t["Estado"] == "Completado" for t in tareas)
    horas = sum(float(e.get("Horas", 0) or 0) for e in edt)
    costo = sum(float(e.get("Costo", 0) or 0) for e in edt)
    sevs = [int(r["Probabilidad"]) * int(r["Impacto"]) for r in riesgos]
    mc_ok = estado_paso("montecarlo")[0]

    def o_pendiente(cond, valor, sub, paso_n):
        return (valor, sub) if cond else ("—", f"Pendiente: paso {paso_n}")

    v1 = o_pendiente(rf or rnf, f"{len(rf) + len(rnf)}", f"{len(rf)} funcionales, {len(rnf)} no funcionales", 1)
    v2 = o_pendiente(tareas, f"{fmt_num(n_ok / len(tareas) * 100) if tareas else 0}%", f"{n_ok} de {len(tareas)} completadas", 4)
    v3 = o_pendiente(edt, f"{fmt_num(horas)} h", fmt_money(costo), 5)
    v4 = o_pendiente(mc_ok, f"{fmt_num(MC_ACTUAL['met']['pt'], 1)} d" if mc_ok else "",
                     MC_ACTUAL["fecha_pt"].strftime("%d/%m/%Y") if mc_ok else "", 6)
    v5 = o_pendiente(riesgos, f"{sum(s >= 12 for s in sevs)}", f"altos de {len(riesgos)} registrados", 7)
    franja([
        {"icono": "layers", "etiqueta": "Requisitos", "valor": v1[0], "sub": v1[1]},
        {"icono": "kanban", "etiqueta": "Avance de tareas", "valor": v2[0], "sub": v2[1]},
        {"icono": "tree", "etiqueta": "Esfuerzo planificado", "valor": v3[0], "sub": v3[1]},
        {"icono": "curve", "etiqueta": "Plazo comprometido", "valor": v4[0], "sub": v4[1], "destacado": mc_ok},
        {"icono": "shield", "etiqueta": "Riesgos altos", "valor": v5[0], "sub": v5[1]},
    ])

    st.write("")
    col_a, col_b = st.columns([1.35, 1], gap="large")
    with col_a:
        seccion("Viabilidad del plan", "Según la simulación Monte Carlo.", "target")
        if mc_ok:
            m, c = MC_ACTUAL["met"], MC_ACTUAL["cfg"]
            html('<div class="mx-panel">'
                 + barra_prob(f"Cumplir el plazo de {fmt_num(c['dias_meta'], 0)} días", m["prob_t"], c["perc"])
                 + barra_prob(f"Cumplir el presupuesto de {fmt_money(c['costo_meta'])}", m["prob_c"], c["perc"])
                 + barra_prob("Cumplir ambas metas", m["prob_ambos"], c["perc"])
                 + f'<div style="color:{T["sub"]};font-size:.8rem;margin-top:6px">La marca vertical es la confianza exigida (P{c["perc"]}).</div></div>')
        else:
            html(f'<div class="mx-panel mx-empty">{ico("curve", 22, T["sub"])}<div>Aparecerá cuando completes el '
                 f'<b>paso 6: Monte Carlo</b>.<br><span>Ingresa tu estimación de días y tus metas de plazo y presupuesto.</span></div></div>')
            st.page_link(PG_MC, label="Ir al paso 6: Monte Carlo", icon=":material/arrow_forward:")

        st.write("")
        seccion("Recomendaciones", "Generadas a partir de lo que llevas.", "flag")
        recos = []
        if siguiente:
            recos.append(("info", T["accent"], f"Tu siguiente paso es <b>{siguiente[0]}. {siguiente[1][1]}</b>: {siguiente[1][3][0].lower() + siguiente[1][3][1:]}"))
        if rf and not rnf:
            recos.append(("alert", T["warn"], "Agrega al menos un requisito no funcional (rendimiento, seguridad, disponibilidad)."))
        if any(s >= 12 for s in sevs):
            recos.append(("alert", T["bad"], f"{sum(s >= 12 for s in sevs)} riesgo(s) de severidad alta necesitan un plan de contingencia."))
        wip = sum(t["Estado"] == "En Proceso" for t in tareas)
        if wip > 3:
            recos.append(("alert", T["warn"], f"Hay {wip} tareas en proceso a la vez. Limitar el WIP a 3 acelera las entregas."))
        if mc_ok:
            m, c = MC_ACTUAL["met"], MC_ACTUAL["cfg"]
            if m["prob_ambos"] < c["perc"]:
                recos.append(("alert", T["warn"], f"La probabilidad de cumplir plazo y presupuesto ({fmt_num(m['prob_ambos'], 1)}%) "
                                                   f"está bajo la confianza exigida ({c['perc']}%). Ajusta metas o alcance."))
            else:
                recos.append(("check", T["ok"], "El plan cumple la confianza exigida en plazo y presupuesto."))
        if not siguiente:
            recos.append(("check", T["ok"], "Completaste todos los pasos. Tus entregables están listos."))
        html('<div class="mx-panel">' + "".join(f'<div class="mx-todo">{ico(i, 18, c)}<span>{t}</span></div>' for i, c, t in recos) + "</div>")

    with col_b:
        seccion("Riesgos principales", "Ordenados por severidad.", "shield")
        if riesgos:
            top = sorted(riesgos, key=lambda r: -int(r["Probabilidad"]) * int(r["Impacto"]))[:5]
            filas = [(r["Riesgo"], f'{int(r["Probabilidad"]) * int(r["Impacto"])} &nbsp;{pill_riesgo(int(r["Probabilidad"]) * int(r["Impacto"]))}')
                     for r in top]
            html(panel("Severidad (probabilidad × impacto)", filas, "alert"))
        else:
            html(f'<div class="mx-panel mx-empty">{ico("shield", 22, T["sub"])}<div>Aún no hay riesgos.<br>'
                 f'<span>Regístralos en el paso 7.</span></div></div>')

        st.write("")
        seccion("Datos del proyecto", "", "users")
        html(panel("Ficha", [
            ("Nombre", ss.nombre_proj),
            ("Equipo", ss.integrantes or "Sin definir"),
            ("Objetivo", (ss.objetivo_text[:60] + "…") if len(ss.objetivo_text) > 60 else (ss.objetivo_text or "Sin definir")),
        ], "info"))


def pagina_requisitos():
    encabezado("layers", "Requisitos", "Define el sistema y sus requisitos funcionales y no funcionales. "
                                        "Estos datos alimentan todos los reportes.")
    seccion("Definición del sistema", icono="info")
    c1, c2 = st.columns(2)
    c1.text_input("Nombre del proyecto", key="nombre_proj")
    c2.text_input("Integrantes / equipo", key="integrantes")
    st.text_area("Objetivo principal", key="objetivo_text", height=80)

    def bloque(titulo, sub, lista_key, prefijo, form_key, placeholder):
        st.write("")
        seccion(titulo, sub, "list")
        with st.form(form_key, clear_on_submit=True, border=False):
            c_in, c_b = st.columns([5, 1])
            nuevo = c_in.text_input("Nuevo", placeholder=placeholder, label_visibility="collapsed")
            if c_b.form_submit_button("Añadir", icon=":material/add:", **ANCHO) and nuevo.strip():
                st.session_state[lista_key].append(nuevo.strip())
                st.rerun()
        lista = st.session_state[lista_key]
        if not lista:
            st.info("Aún no hay requisitos. Escribe uno arriba y pulsa Añadir.")
            return
        borrar = None
        v = version_lista(lista_key)
        for i, item in enumerate(lista):
            c_tag, c_input, c_del = st.columns([0.7, 6, 0.45], vertical_alignment="center")
            c_tag.markdown(f"**{prefijo}-{i + 1:02d}**")
            lista[i] = c_input.text_input(f"{prefijo}_{i}", value=item, label_visibility="collapsed",
                                          key=f"{prefijo.lower()}_field_{v}_{i}")
            if c_del.button("", icon=":material/delete:", key=f"del_{prefijo.lower()}_{v}_{i}", help="Eliminar requisito"):
                borrar = i
        if borrar is not None:
            eliminar_fila(lista_key, borrar, f"{prefijo.lower()}_field_", f"del_{prefijo.lower()}_")
            st.rerun()

    col_rf, col_rnf = st.columns(2, gap="large")
    with col_rf:
        bloque("Requisitos funcionales", "Lo que el sistema debe hacer.", "rf_list", "RF", "form_add_rf",
               "Escribe un requisito funcional")
    with col_rnf:
        bloque("Requisitos no funcionales", "Cómo debe comportarse: rendimiento, seguridad, disponibilidad.",
               "rnf_list", "RNF", "form_add_rnf", "Escribe un requisito no funcional")

    st.divider()
    seccion("Exportar", icono="archive")
    d1, d2, _ = st.columns([1, 1, 1])
    d1.download_button("Requisitos en Excel", generar_excel_requisitos(), f"Requisitos_{slug(st.session_state.nombre_proj)}.xlsx",
                       XLSX, icon=":material/table_view:", **ANCHO)
    d2.download_button("Documento ERS en Word", generar_word_ers(), f"ERS_Requisitos_{slug(st.session_state.nombre_proj)}.docx",
                       DOCX, icon=":material/description:", **ANCHO)


def pagina_modelado():
    encabezado("nodes", "Modelado", "Arquitectura visual con sintaxis Mermaid. Edita a la izquierda y revisa el diagrama a la derecha.")
    c_code, c_prev = st.columns([1, 1.4], gap="large")
    with c_code:
        st.text_area("Sintaxis Mermaid", key="modelo_mermaid", height=300)
        st.download_button("Descargar diagrama (.mmd)", st.session_state.modelo_mermaid.encode("utf-8"),
                           f"Diagrama_Modelado_{slug(st.session_state.nombre_proj)}.mmd", "text/plain",
                           icon=":material/download:", **ANCHO)
    with c_prev:
        mostrar_mermaid(st.session_state.modelo_mermaid, "380px", key=f"mm_modelo_{st.session_state.tema}")


def pagina_flujos():
    encabezado("flow", "Flujos y secuencias", "Parte de una plantilla y adáptala; la vista previa se actualiza al instante.")
    c1, c2 = st.columns([1, 1.4], gap="large")
    with c1:
        st.selectbox("Plantilla", options=list(PLANTILLAS_FLUJO.keys()), key="select_tipo_flujo",
                     on_change=cambiar_plantilla_flujo)
        st.text_area("Código del diagrama", key="flujo_codigo", height=260)
        st.download_button("Descargar diagrama de flujo (.mmd)", st.session_state.flujo_codigo.encode("utf-8"),
                           f"Flujo_Secuencia_{slug(st.session_state.nombre_proj)}.mmd", "text/plain",
                           icon=":material/download:", **ANCHO)
    with c2:
        mostrar_mermaid(st.session_state.flujo_codigo, "420px", key=f"mm_flujo_{st.session_state.tema}")


def pagina_gestion():
    encabezado("kanban", "Tablero Kanban", "Registra tareas y muévelas cambiando su estado.")
    tareas = st.session_state.kanban_tasks

    with st.form("form_kanban", clear_on_submit=True, border=False):
        c_t, c_s, c_b = st.columns([4, 2, 1])
        nueva = c_t.text_input("Tarea", placeholder="Nueva tarea", label_visibility="collapsed")
        estado = c_s.selectbox("Estado", ESTADOS_KANBAN, label_visibility="collapsed")
        if c_b.form_submit_button("Añadir", icon=":material/add:", **ANCHO) and nueva.strip():
            tareas.append({"Tarea": nueva.strip(), "Estado": estado})
            st.rerun()

    if not tareas:
        st.info("No hay tareas. Escribe la primera arriba para empezar el tablero.")
    else:
        n_ok = sum(t["Estado"] == "Completado" for t in tareas)
        wip = sum(t["Estado"] == "En Proceso" for t in tareas)
        franja([
            {"icono": "list", "etiqueta": "Tareas", "valor": len(tareas)},
            {"icono": "progress", "etiqueta": "En proceso", "valor": wip, "sub": "Límite sugerido: 3"},
            {"icono": "check", "etiqueta": "Completadas", "valor": n_ok, "sub": f"{fmt_num(n_ok / len(tareas) * 100)}% de avance", "destacado": True},
        ])
        st.write("")
        if wip > 3:
            st.warning(f"Hay {wip} tareas en proceso a la vez. Limitar el trabajo en curso ayuda a terminar antes.")
        estilos = {"Pendiente": (T["sub"], "circle"), "En Proceso": (T["warn"], "progress"), "Completado": (T["ok"], "check")}
        cols = st.columns(3)
        for n_col, (col, est) in enumerate(zip(cols, ESTADOS_KANBAN)):
            color, icono = estilos[est]
            items = [t for t in tareas if t["Estado"] == est]
            cuerpo = "".join(f'<div class="it">{t["Tarea"]}</div>' for t in items) or '<div class="empty">Sin tareas</div>'
            col.markdown(f'<div class="mx-kb" style="--c:{color};--d:{n_col * 0.07:.2f}s"><div class="h">'
                         f'<span class="l">{ico(icono, 16)}{est}</span><span class="n">{len(items)}</span></div>{cuerpo}</div>',
                         unsafe_allow_html=True)

        st.write("")
        with st.expander("Editar o eliminar tareas", icon=":material/edit:"):
            borrar = None
            v = version_lista("kanban_tasks")
            for i, t in enumerate(tareas):
                c_txt, c_sel, c_d = st.columns([4, 2, 0.45], vertical_alignment="center")
                t["Tarea"] = c_txt.text_input(f"kt_{i}", value=t["Tarea"], label_visibility="collapsed", key=f"kt_in_{v}_{i}")
                t["Estado"] = c_sel.selectbox(f"ks_{i}", ESTADOS_KANBAN, index=ESTADOS_KANBAN.index(t["Estado"]),
                                              label_visibility="collapsed", key=f"ks_in_{v}_{i}")
                if c_d.button("", icon=":material/delete:", key=f"del_k_{v}_{i}", help="Eliminar tarea"):
                    borrar = i
            if borrar is not None:
                eliminar_fila("kanban_tasks", borrar, "kt_in_", "ks_in_", "del_k_")
                st.rerun()

    st.divider()
    seccion("Exportar", icono="archive")
    d1, d2, _ = st.columns([1, 1, 1])
    d1.download_button("Kanban en Excel", generar_excel_kanban(), f"Kanban_{slug(st.session_state.nombre_proj)}.xlsx",
                       XLSX, icon=":material/table_view:", **ANCHO)
    d2.download_button("Tareas en CSV", pd.DataFrame(tareas).to_csv(index=False).encode("utf-8"),
                       f"Kanban_{slug(st.session_state.nombre_proj)}.csv", "text/csv", icon=":material/download:", **ANCHO)


def pagina_edt():
    encabezado("tree", "Estructura de desglose de trabajo",
               "Proyecto, fases, subfases y paquetes de trabajo con horas, costo y documentos. "
               "Los paquetes también pueden alimentar la simulación Monte Carlo.")
    edt = st.session_state.edt_list
    df = pd.DataFrame(edt)
    tot_h = float(df["Horas"].sum()) if not df.empty else 0
    tot_c = float(df["Costo"].sum()) if not df.empty else 0
    franja([
        {"icono": "clock", "etiqueta": "Horas totales", "valor": f"{fmt_num(tot_h)} h", "sub": f"≈ {fmt_num(tot_h / 8, 1)} días de 8 h"},
        {"icono": "wallet", "etiqueta": "Costo estimado", "valor": fmt_money(tot_c), "destacado": True},
        {"icono": "list", "etiqueta": "Paquetes de trabajo", "valor": len(df)},
        {"icono": "folder", "etiqueta": "Fases", "valor": df["Fase"].nunique() if not df.empty else 0},
    ])

    st.write("")
    t_arbol, t_dist = st.tabs(["Diagrama jerárquico", "Distribución por fase"])
    with t_arbol:
        if edt:
            mostrar_mermaid(generar_mermaid_edt(edt, st.session_state.nombre_proj), "470px",
                            key=f"mm_edt_{st.session_state.tema}")
        else:
            st.info("Agrega un paquete de trabajo para ver el árbol.")
    with t_dist:
        if not df.empty:
            metrica = st.segmented_control("Métrica", ["Horas", "Costo"], key="rad_edt_metric")
            metrica = metrica or "Horas"
            df_f = df.groupby("Fase", as_index=False)[["Horas", "Costo"]].sum().sort_values(metrica)
            fig = go.Figure(go.Bar(
                x=df_f[metrica], y=df_f["Fase"], orientation="h",
                marker=dict(color=T["accent"], line=dict(width=0)),
                text=[fmt_num(v) + (" h" if metrica == "Horas" else " CLP") for v in df_f[metrica]],
                textposition="outside", textfont=dict(color=T["text"]),
                hovertemplate="%{y}<br>%{x:,.0f}<extra></extra>"))
            estilo_fig(fig, height=max(240, 90 * len(df_f)), showlegend=False)
            fig.update_xaxes(showgrid=False, showticklabels=False)
            mostrar_fig(fig, key="edt_bar")
        else:
            st.info("Agrega datos para ver la distribución.")
    st.divider()
    if True:
        with st.form("form_add_edt", clear_on_submit=True):
            seccion("Nuevo paquete de trabajo", icono="tree")
            c1, c2, c3 = st.columns(3)
            fase_in = c1.text_input("Fase principal", placeholder="Ej: Fase 1: Preparación")
            sub_in = c2.text_input("Subfase / grupo (opcional)", placeholder="Ej: Pintura aerosol")
            paq_in = c3.text_input("Paquete de trabajo", placeholder="Ej: Pintar arbusto")
            c4, c5, c6 = st.columns(3)
            horas_in = c4.number_input("Horas estimadas", min_value=0, value=4, step=1)
            costo_in = c5.number_input("Costo estimado ($)", min_value=0, value=15000, step=1000)
            doc_in = c6.text_input("Documentos / referencias", placeholder="Ej: Ficha técnica")
            if st.form_submit_button("Añadir paquete", icon=":material/add:"):
                if fase_in.strip() and paq_in.strip():
                    edt.append({"Fase": fase_in.strip(), "Subfase": sub_in.strip(), "Paquete": paq_in.strip(),
                                "Horas": float(horas_in), "Costo": float(costo_in), "Documentos": doc_in.strip()})
                    st.rerun()
                else:
                    st.error("Completa al menos la fase principal y el paquete de trabajo.")

        st.write("")
        seccion("Editar paquetes", "Los cambios se aplican al instante en el árbol, los totales y Monte Carlo.", "list")
        if edt:
            anchos = [2, 1.8, 2.2, 0.9, 1.1, 1.8, 0.45]
            cab = st.columns(anchos)
            for c, t in zip(cab, ["Fase", "Subfase", "Paquete", "Horas", "Costo ($)", "Documentos", ""]):
                c.caption(t)
            borrar = None
            v = version_lista("edt_list")
            for i, item in enumerate(edt):
                cs = st.columns(anchos, vertical_alignment="center")
                item["Fase"] = cs[0].text_input(f"f_{i}", value=item.get("Fase", ""), key=f"edt_fase_{v}_{i}", label_visibility="collapsed")
                item["Subfase"] = cs[1].text_input(f"sf_{i}", value=item.get("Subfase", ""), key=f"edt_sfase_{v}_{i}", label_visibility="collapsed")
                item["Paquete"] = cs[2].text_input(f"p_{i}", value=item.get("Paquete", ""), key=f"edt_pkg_{v}_{i}", label_visibility="collapsed")
                item["Horas"] = cs[3].number_input(f"h_{i}", value=float(item.get("Horas", 0)), min_value=0.0, step=1.0,
                                                   key=f"edt_hrs_{v}_{i}", label_visibility="collapsed")
                item["Costo"] = cs[4].number_input(f"c_{i}", value=float(item.get("Costo", 0)), min_value=0.0, step=1000.0,
                                                   key=f"edt_cost_{v}_{i}", label_visibility="collapsed")
                item["Documentos"] = cs[5].text_input(f"d_{i}", value=item.get("Documentos", ""), key=f"edt_doc_{v}_{i}", label_visibility="collapsed")
                if cs[6].button("", icon=":material/delete:", key=f"del_edt_{v}_{i}", help="Eliminar paquete"):
                    borrar = i
            if borrar is not None:
                eliminar_fila("edt_list", borrar, "edt_fase_", "edt_sfase_", "edt_pkg_", "edt_hrs_", "edt_cost_", "edt_doc_", "del_edt_")
                st.rerun()

    st.divider()
    seccion("Exportar", icono="archive")
    d1, d2, _ = st.columns([1, 1, 1])
    d1.download_button("EDT en Excel", generar_excel_edt(), f"EDT_{slug(st.session_state.nombre_proj)}.xlsx",
                       XLSX, icon=":material/table_view:", **ANCHO)
    d2.download_button("EDT en CSV", pd.DataFrame(edt).to_csv(index=False).encode("utf-8"),
                       f"EDT_{slug(st.session_state.nombre_proj)}.csv", "text/csv", icon=":material/download:", **ANCHO)


def pagina_montecarlo():
    encabezado("curve", "Simulación Monte Carlo",
               "Simula miles de versiones posibles del proyecto para saber qué probabilidad hay de cumplir "
               "el plazo y el presupuesto, y qué compromiso conviene asumir.")

    with st.expander("Cómo funciona y cómo leer los resultados", icon=":material/help:"):
        st.markdown("""
**La idea.** En vez de una sola estimación ("tardaremos 20 días"), defines un rango de incertidumbre. El motor sortea
miles de duraciones dentro de ese rango, como si ejecutaras el proyecto miles de veces, y cuenta cuántas veces se cumplen tus metas.

**Distribuciones**
- **Beta-PERT** (recomendada, estándar PMBOK): usa optimista, más probable y pesimista, y concentra los resultados cerca del valor más probable.
- **Triangular**: los mismos tres valores, con más peso en los extremos. Es más conservadora.
- **Normal**: procesos conocidos, definidos por promedio y desviación.
- **Uniforme**: cuando solo conoces el mínimo y el máximo.

**Percentiles.** P50 es una apuesta 50/50. P85 (o el que elijas) es el plazo que puedes prometer con esa confianza. P95 cubre casi todos los imprevistos.

**Reserva de contingencia.** La diferencia entre el P50 y el percentil de compromiso: el colchón de tiempo y dinero que conviene negociar.
        """)

    ss = st.session_state
    col_cfg, col_res = st.columns([1, 2.05], gap="large")

    # ---------------- Panel de parámetros ----------------
    with col_cfg:
        with st.container(border=True):
            paso(1, "Estimación de la duración")
            st.radio("Fuente", ["Estimación manual (3 puntos)", "Desde la EDT (por paquete)"], key="mc_fuente",
                     help="La opción EDT simula cada paquete de trabajo por separado y suma sus duraciones.")
            st.selectbox("Distribución", DISTRIBUCIONES, key="mc_dist_type")
            if str(ss.mc_fuente).startswith("Desde"):
                st.number_input("Horas de trabajo por día", min_value=1.0, max_value=24.0, step=0.5, format="%.1f", key="mc_horas_dia")
                c1, c2 = st.columns(2)
                c1.slider("Mejor caso (−%)", 0, 80, step=5, key="mc_pct_opt")
                c2.slider("Peor caso (+%)", 0, 300, step=10, key="mc_pct_pes")
                base_h = sum(float(e.get("Horas", 0) or 0) for e in ss.edt_list)
                st.caption(f"Base EDT: {len(ss.edt_list)} paquetes, {fmt_num(base_h, 1)} h ≈ "
                           f"{fmt_num(base_h / max(float(ss.mc_horas_dia), 1), 1)} días si se ejecutan en secuencia.")
            elif ss.mc_dist_type in ("Triangular", "Beta-PERT"):
                st.number_input("Optimista (días)", min_value=0.5, step=1.0, format="%.1f", key="mc_opt")
                st.number_input("Más probable (días)", min_value=0.5, step=1.0, format="%.1f", key="mc_prob")
                st.number_input("Pesimista (días)", min_value=0.5, step=1.0, format="%.1f", key="mc_pes")
                st.caption(f"PERT clásico (O + 4M + P) / 6 = {fmt_num((ss.mc_opt + 4 * ss.mc_prob + ss.mc_pes) / 6, 1)} días")
            elif ss.mc_dist_type == "Normal":
                st.number_input("Duración promedio (días)", min_value=1.0, step=1.0, format="%.1f", key="mc_mean")
                st.number_input("Desviación estándar (días)", min_value=0.1, step=0.5, format="%.1f", key="mc_std")
                st.caption(f"El 95% de los escenarios caerá entre {fmt_num(max(ss.mc_mean - 2 * ss.mc_std, 0.5), 1)} "
                           f"y {fmt_num(ss.mc_mean + 2 * ss.mc_std, 1)} días.")
            else:
                st.number_input("Mínimo (días)", min_value=0.5, step=1.0, format="%.1f", key="mc_min_u")
                st.number_input("Máximo (días)", min_value=0.5, step=1.0, format="%.1f", key="mc_max_u")

            paso(2, "Costos")
            st.number_input("Costo fijo base ($)", min_value=0, step=10000, key="mc_cfijo")
            if str(ss.mc_fuente).startswith("Desde"):
                st.caption("El costo variable sale de cada paquete de la EDT y crece si el paquete se demora.")
            else:
                st.number_input("Costo operativo por día ($)", min_value=0, step=1000, key="mc_cdiario")
                st.slider("Variabilidad del costo diario (±%)", 0, 30, key="mc_var_costo")

            paso(3, "Metas")
            st.number_input("Plazo meta (días)", min_value=0.5, step=1.0, format="%.1f", key="mc_dias_meta")
            st.number_input("Presupuesto meta ($)", min_value=0, step=25000, key="mc_costo_meta")
            st.slider("Confianza exigida (percentil)", 50, 99, key="mc_target_perc",
                      help="Con qué probabilidad quieres poder cumplir lo que prometas.")

            paso(4, "Calendario y motor")
            st.date_input("Fecha de inicio", key="mc_f_inicio", format="DD/MM/YYYY")
            st.toggle("Contar solo días hábiles", key="mc_habiles")
            st.select_slider("Iteraciones", options=[500, 1000, 2000, 5000, 10000, 20000, 50000], key="mc_sims")
            st.toggle("Resultados reproducibles", key="mc_seed", help="Usa siempre la misma semilla aleatoria.")
            if not ss.mc_seed:
                st.button("Nueva corrida aleatoria", icon=":material/casino:", on_click=nuevas_semilla, **ANCHO)

    # ---------------- Resultados ----------------
    with col_res:
        if MC_ACTUAL is None:
            st.error(MC_ERROR, icon=":material/error:")
            return
        res, met, cfg = MC_ACTUAL["res"], MC_ACTUAL["met"], MC_ACTUAL["cfg"]
        dias, costos = res["dias"], res["costos"]
        perc, dias_meta, costo_meta = cfg["perc"], cfg["dias_meta"], cfg["costo_meta"]
        inicio, habiles = ss.mc_f_inicio, cfg["habiles"]
        ok_t, ok_c = dias <= dias_meta, costos <= costo_meta

        nivel, titulo_v = veredicto(met, perc)
        color_v = T[nivel]
        n_ok = int(round(met["prob_ambos"]))
        n_mid = max(0, min(int(round(met["prob_ambos"] + met["prob_uno"])) - n_ok, 100 - n_ok))
        clases = ["ok"] * n_ok + ["mid"] * n_mid + ["bad"] * (100 - n_ok - n_mid)
        puntos = "".join(f'<i class="{c}" style="animation-delay:{i * 0.006:.3f}s"></i>' for i, c in enumerate(clases))
        if nivel == "ok":
            reco = (f"Puedes comprometer <b>{fmt_num(met['pt'], 1)} días</b> (término el <b>{MC_ACTUAL['fecha_pt'].strftime('%d/%m/%Y')}</b>) "
                    f"y <b>{fmt_money(met['ptc'])}</b> con un {perc}% de confianza.")
        else:
            falta = []
            if met["pt"] > dias_meta:
                falta.append(f"ampliar el plazo a <b>{fmt_num(np.ceil(met['pt']))} días</b>")
            if met["ptc"] > costo_meta:
                falta.append(f"ampliar el presupuesto a <b>{fmt_money(met['ptc'])}</b>")
            reco = (f"Para llegar a un {perc}% de confianza habría que " + " y ".join(falta) + ", o reducir el alcance."
                    if falta else "Cada meta se cumple por separado, pero no ambas a la vez: revisa la holgura.")
        icono_v = {"ok": "check", "warn": "alert", "bad": "cross"}[nivel]
        html(f'<div class="mx-verdict" style="--c:{color_v}"><div>'
             f'<span class="tag">{ico(icono_v, 16)}{titulo_v}</span>'
             f'<div class="big">De cada 100 escenarios, <b>{n_ok}</b> terminan a tiempo y dentro del presupuesto.</div>'
             f'<div class="txt">{reco} Reserva de contingencia sugerida: <b>+{fmt_num(met["reserva_dias"], 1)} días</b> '
             f'y <b>+{fmt_money(met["reserva_costo"])}</b> sobre la mediana.</div></div>'
             f'<div><div class="mx-dots">{puntos}</div><div class="mx-leg">'
             f'<span style="--c:{T["ok"]}">Ambas metas</span><span style="--c:{T["warn"]}">Una</span>'
             f'<span style="--c:{T["bad"]}">Ninguna</span></div></div></div>')

        st.write("")
        franja([
            {"icono": "calendar", "etiqueta": "Entrega probable (P50)", "valor": f"{fmt_num(met['p50'], 1)} d",
             "sub": MC_ACTUAL["fecha_p50"].strftime("%d/%m/%Y")},
            {"icono": "target", "etiqueta": f"Plazo a comprometer (P{perc})", "valor": f"{fmt_num(met['pt'], 1)} d",
             "sub": MC_ACTUAL["fecha_pt"].strftime("%d/%m/%Y"), "destacado": True},
            {"icono": "alert", "etiqueta": "Pesimista (P95)", "valor": f"{fmt_num(met['p95'], 1)} d",
             "sub": MC_ACTUAL["fecha_p95"].strftime("%d/%m/%Y")},
            {"icono": "wallet", "etiqueta": f"Presupuesto P{perc}", "valor": fmt_money(met["ptc"]),
             "sub": f"Mediana {fmt_money(met['p50c'])}"},
        ])

        st.write("")
        html('<div class="mx-panel">'
             + barra_prob(f"Probabilidad de cumplir el plazo ({fmt_num(dias_meta, 0)} días)", met["prob_t"], perc)
             + barra_prob(f"Probabilidad de cumplir el presupuesto ({fmt_money(costo_meta)})", met["prob_c"], perc)
             + barra_prob("Probabilidad de cumplir ambas metas", met["prob_ambos"], perc)
             + f'<div style="color:{T["sub"]};font-size:.8rem;margin-top:6px">La marca vertical indica la confianza exigida (P{perc}).</div></div>')

        st.write("")
        nombres = ["Plazo", "Curva S", "Plazo vs. costo", "Costo", "Convergencia"]
        if res["componentes"] is not None:
            nombres.append("Sensibilidad")
        tabs = st.tabs(nombres)

        def marcar(fig, marcas):
            for i, (x, txt, col, dash) in enumerate(marcas):
                fig.add_vline(x=x, line_color=col, line_dash=dash, line_width=1.6)
                fig.add_annotation(x=x, y=1 - i * 0.08, yref="paper", text=txt, showarrow=False, xanchor="left",
                                   xshift=5, font=dict(color=col, size=11.5), bgcolor=rgba(T["card"], 0.85))

        def histograma(datos, meta, titulo_x, dinero=False):
            conteos, bordes = np.histogram(datos, bins=42)
            centros = (bordes[:-1] + bordes[1:]) / 2
            pct = conteos / len(datos) * 100
            dentro = bordes[1:] <= meta
            fmt_h = "$%{customdata[0]:,.0f} – $%{customdata[1]:,.0f}" if dinero else "%{customdata[0]:.1f} – %{customdata[1]:.1f} días"
            fig = go.Figure()
            for mask, nombre, col in [(dentro, "Cumple la meta", T["ok"]), (~dentro, "Excede la meta", T["bad"])]:
                if mask.any():
                    fig.add_bar(x=centros[mask], y=pct[mask], width=np.diff(bordes)[mask] * 0.92, name=nombre,
                                marker=dict(color=rgba(col, 0.72), line=dict(width=0)),
                                customdata=np.stack([bordes[:-1][mask], bordes[1:][mask]], axis=-1),
                                hovertemplate=fmt_h + "<br>%{y:.2f}% de los escenarios<extra></extra>")
            fig.update_layout(barmode="overlay", bargap=0)
            estilo_fig(fig, height=390)
            fig.update_xaxes(title=titulo_x)
            fig.update_yaxes(title="% de escenarios", ticksuffix="%")
            return fig

        with tabs[0]:
            fig = histograma(dias, dias_meta, "Duración del proyecto (días)")
            marcar(fig, [(dias_meta, f"Meta {fmt_num(dias_meta, 1)}", T["text"], "dash"),
                         (met["p50"], f"P50 {fmt_num(met['p50'], 1)}", T["sub"], "dot"),
                         (met["pt"], f"P{perc} {fmt_num(met['pt'], 1)}", T["accent"], "solid"),
                         (met["p95"], f"P95 {fmt_num(met['p95'], 1)}", T["warn"], "dot")])
            mostrar_fig(fig, key="mc_hist")
            st.caption(f"Verde: {fmt_num(met['prob_t'], 1)}% de los escenarios termina dentro del plazo. "
                       f"Rojo: {fmt_num(met['riesgo_t'], 1)}% se atrasa. Promedio {fmt_num(met['media'], 1)} días, "
                       f"desviación {fmt_num(met['desv'], 1)} días.")

        with tabs[1]:
            xs = np.sort(dias)
            ys = np.arange(1, len(xs) + 1) / len(xs) * 100
            idx = np.linspace(0, len(xs) - 1, min(len(xs), 700)).astype(int)
            fig = go.Figure()
            fig.add_scatter(x=xs[idx], y=ys[idx], mode="lines", line=dict(color=T["accent"], width=2.6),
                            fill="tozeroy", fillcolor=rgba(T["accent"], 0.08),
                            hovertemplate="Terminar en ≤ %{x:.1f} días<br>Probabilidad %{y:.1f}%<extra></extra>")
            pts = [(met["p50"], 50, "P50", T["sub"]), (met["pt"], perc, f"P{perc}", T["accent"]), (met["p95"], 95, "P95", T["warn"])]
            fig.add_scatter(x=[p[0] for p in pts], y=[p[1] for p in pts], mode="markers+text",
                            text=[f"{p[2]}  {fmt_num(p[0], 1)} d" for p in pts], textposition="middle left",
                            marker=dict(size=10, color=[p[3] for p in pts], line=dict(width=2, color=T["card"])),
                            textfont=dict(color=T["text"], size=11.5), hoverinfo="skip")
            fig.add_shape(type="line", x0=dias_meta, x1=dias_meta, y0=0, y1=met["prob_t"], line=dict(color=T["text"], dash="dash", width=1.2))
            fig.add_shape(type="line", x0=xs.min(), x1=dias_meta, y0=met["prob_t"], y1=met["prob_t"], line=dict(color=T["text"], dash="dash", width=1.2))
            fig.add_annotation(x=dias_meta, y=met["prob_t"], text=f"Meta: {fmt_num(met['prob_t'], 1)}%", showarrow=True,
                               arrowhead=0, ax=55, ay=28, font=dict(color=T["text"]), bgcolor=rgba(T["card"], 0.9),
                               bordercolor=T["border"])
            estilo_fig(fig, height=390, showlegend=False)
            fig.update_xaxes(title="Duración (días)")
            fig.update_yaxes(title="Probabilidad de terminar en ese plazo o antes", ticksuffix="%", range=[0, 103])
            mostrar_fig(fig, key="mc_cdf")
            st.caption("Elige un número de días en el eje horizontal y sube hasta la curva: la altura es la probabilidad de cumplirlo.")

            q1, q2 = st.columns(2)
            with q1:
                x_d = st.number_input("Si me comprometo a … días", min_value=0.5, value=float(np.ceil(met["pt"])),
                                      step=1.0, format="%.1f", key="mc_calc_dias")
                html(f'<div class="mx-answer">Probabilidad de cumplir: <b>{fmt_num((dias <= x_d).mean() * 100, 1)}%</b><br>'
                     f'Término el {fecha_fin(inicio, x_d, habiles).strftime("%d/%m/%Y")}</div>')
            with q2:
                conf = st.slider("Quiero esta confianza (%)", 50, 99, 90, key="mc_calc_conf")
                d_c, c_c = float(np.percentile(dias, conf)), float(np.percentile(costos, conf))
                html(f'<div class="mx-answer">Necesitas <b>{fmt_num(d_c, 1)} días</b> y <b>{fmt_money(c_c)}</b><br>'
                     f'Término el {fecha_fin(inicio, d_c, habiles).strftime("%d/%m/%Y")}</div>')

        with tabs[2]:
            cat = np.where(ok_t & ok_c, "Cumple ambas", np.where(ok_t, "Solo plazo", np.where(ok_c, "Solo presupuesto", "Ninguna")))
            muestra = np.random.default_rng(7).choice(len(dias), size=min(2500, len(dias)), replace=False)
            fig = go.Figure()
            for nombre_c, col in {"Cumple ambas": T["ok"], "Solo plazo": T["accent2"], "Solo presupuesto": T["warn"], "Ninguna": T["bad"]}.items():
                m_all = cat == nombre_c
                if not m_all.any():
                    continue
                m = m_all[muestra]
                fig.add_trace(go.Scattergl(x=dias[muestra][m], y=costos[muestra][m], mode="markers",
                                           name=f"{nombre_c} ({fmt_num(m_all.mean() * 100, 1)}%)",
                                           marker=dict(color=rgba(col, 0.55), size=5.5),
                                           hovertemplate="%{x:.1f} días<br>$%{y:,.0f}<extra></extra>"))
            fig.add_vline(x=dias_meta, line_dash="dash", line_color=T["text"], line_width=1.2)
            fig.add_hline(y=costo_meta, line_dash="dash", line_color=T["text"], line_width=1.2)
            estilo_fig(fig, height=410)
            fig.update_xaxes(title="Duración (días)")
            fig.update_yaxes(title="Costo total ($)", tickprefix="$")
            mostrar_fig(fig, key="mc_scatter")
            st.caption("Cada punto es un escenario. El cuadrante inferior izquierdo de las líneas punteadas cumple ambas metas.")

        with tabs[3]:
            fig = histograma(costos, costo_meta, "Costo total ($)", dinero=True)
            marcar(fig, [(costo_meta, f"Meta {fmt_money(costo_meta)}", T["text"], "dash"),
                         (met["p50c"], f"P50 {fmt_money(met['p50c'])}", T["sub"], "dot"),
                         (met["ptc"], f"P{perc} {fmt_money(met['ptc'])}", T["accent"], "solid")])
            fig.update_xaxes(tickprefix="$")
            mostrar_fig(fig, key="mc_cost")
            st.caption(f"{fmt_num(met['prob_c'], 1)}% de los escenarios queda dentro del presupuesto. Promedio {fmt_money(met['media_c'])}.")

        with tabs[4]:
            cps = np.unique(np.linspace(min(100, len(dias)), len(dias), 70).astype(int))
            media_run = np.cumsum(dias)[cps - 1] / cps
            p_run = np.array([np.percentile(dias[:c], perc) for c in cps])
            fig = go.Figure()
            fig.add_hrect(y0=met["pt"] * 0.99, y1=met["pt"] * 1.01, fillcolor=rgba(T["accent"], 0.10), line_width=0)
            fig.add_scatter(x=cps, y=p_run, mode="lines", name=f"P{perc} acumulado", line=dict(color=T["accent"], width=2.4))
            fig.add_scatter(x=cps, y=media_run, mode="lines", name="Promedio acumulado", line=dict(color=T["sub"], width=1.8))
            estilo_fig(fig, height=360)
            fig.update_xaxes(title="Iteraciones ejecutadas")
            fig.update_yaxes(title="Días")
            mostrar_fig(fig, key="mc_conv")
            estable = bool(np.all(np.abs(p_run[len(p_run) // 2:] - met["pt"]) <= met["pt"] * 0.01))
            st.caption(("Resultado estable: " if estable else "Resultado aún variable: ")
                       + f"en la segunda mitad de la simulación el P{perc} se mueve "
                       + ("menos de ±1% (banda sombreada)." if estable else "más de ±1%. Aumenta las iteraciones."))

        if res["componentes"] is not None:
            with tabs[5]:
                comp = res["componentes"]
                var_total = float(np.var(dias))
                aporte = np.array([np.cov(comp[:, j], dias)[0, 1] / var_total * 100 if var_total > 0 else 0
                                   for j in range(comp.shape[1])])
                df_s = pd.DataFrame({"Paquete": res["nombres"], "Aporte": aporte}).sort_values("Aporte").tail(12)
                fig = go.Figure(go.Bar(
                    x=df_s["Aporte"], y=df_s["Paquete"], orientation="h",
                    marker_color=[T["accent"] if v == df_s["Aporte"].max() else rgba(T["sub"], 0.55) for v in df_s["Aporte"]],
                    text=[f"{fmt_num(v, 1)}%" for v in df_s["Aporte"]], textposition="outside", textfont=dict(color=T["text"]),
                    hovertemplate="%{y}<br>Aporte a la incertidumbre: %{x:.1f}%<extra></extra>"))
                estilo_fig(fig, height=max(260, 56 * len(df_s)), showlegend=False)
                fig.update_xaxes(title="Aporte a la incertidumbre total (%)", ticksuffix="%")
                mostrar_fig(fig, key="mc_sens")
                top = df_s.iloc[-1]
                st.caption(f"'{top['Paquete']}' explica el {fmt_num(top['Aporte'], 1)}% de la variación del plazo. "
                           "Afinar su estimación es lo que más reduce la incertidumbre.")

    # ---------------- Detalle (ancho completo) ----------------
    st.write("")
    seccion("Tabla de confianza", "Plazo, fecha y presupuesto para cada nivel de confianza.", "list")
    tp = tabla_percentiles(dias, costos, inicio, habiles, perc)
    df_tp = pd.DataFrame([{
        "Confianza": r["Confianza"], "Días": fmt_num(r["Días"], 1), "Fecha de término": r["Fecha"].strftime("%d/%m/%Y"),
        "Presupuesto": fmt_money(r["Costo"]),
        "Lectura": f"{r['p']}% de los escenarios termina en {fmt_num(r['Días'], 1)} días o menos",
    } for r in tp])
    html(tabla_html(df_tp, resaltar=lambda f: f[0] == f"P{perc}"))

    st.write("")
    c1, c2 = st.columns(2, gap="large")
    with c1:
        seccion("Inspeccionar un escenario", "Revisa una iteración concreta del sorteo.", "search")
        atajo = st.segmented_control("Ir a", ["Número", "Más rápido", "Mediano", "Más lento"], default="Número",
                                     key="mc_atajo", label_visibility="collapsed")
        if atajo in (None, "Número"):
            it = st.slider("Iteración", 1, len(dias), min(100, len(dias)), key="mc_slider_iter")
        else:
            orden = np.argsort(dias)
            it = int({"Más rápido": orden[0], "Mediano": orden[len(orden) // 2], "Más lento": orden[-1]}[atajo]) + 1
        d_it, c_it = float(dias[it - 1]), float(costos[it - 1])
        html(panel(f"Escenario {fmt_num(it)} de {fmt_num(len(dias))}", [
            ("Duración", f"{fmt_num(d_it, 1)} días"),
            ("Costo", fmt_money(c_it)),
            ("Fecha de término", fecha_fin(inicio, d_it, habiles).strftime("%d/%m/%Y")),
            ("Posición", f"más lento que el {fmt_num((dias <= d_it).mean() * 100)}% de los escenarios"),
            ("Plazo", pill("Cumple", "ok", "check") if d_it <= dias_meta else pill("Atrasado", "bad", "cross")),
            ("Presupuesto", pill("Dentro", "ok", "check") if c_it <= costo_meta else pill("Sobrecosto", "bad", "cross")),
        ], "search"))
    with c2:
        seccion("Capacidad anual del equipo", "Cuántos proyectos de este tamaño caben en un año.", "repeat")
        dias_ano = 250 if habiles else 365
        cap50 = dias_ano / met["p50"] if met["p50"] > 0 else 0
        capt = dias_ano / met["pt"] if met["pt"] > 0 else 0
        html(panel(f"Base de {dias_ano} {'días hábiles' if habiles else 'días corridos'} por año", [
            ("Capacidad mediana (P50)", f"{fmt_num(cap50, 1)} proyectos/año"),
            (f"Capacidad segura (P{perc})", f"{fmt_num(capt, 1)} proyectos/año"),
            ("Proyectos completos asegurados", f"{int(capt)} por año"),
            (f"Riesgo asumido al planificar con P{perc}", f"{100 - perc}%"),
        ], "repeat"))

    df_sims = pd.DataFrame({
        "Iteración": np.arange(1, len(dias) + 1), "Duración (días)": np.round(dias, 1),
        "Costo ($)": np.round(costos, 0), "Cumple plazo": ok_t, "Cumple presupuesto": ok_c,
    })
    st.write("")
    with st.expander("Escenarios extremos", icon=":material/swap_vert:"):
        def fmt_ext(df):
            return pd.DataFrame({
                "Iteración": df["Iteración"].map(fmt_num),
                "Duración": df["Duración (días)"].map(lambda v: f"{fmt_num(v, 1)} d"),
                "Costo": df["Costo ($)"].map(fmt_money),
                "Plazo": df["Cumple plazo"].map(lambda v: pill("Cumple", "ok") if v else pill("Atraso", "bad")),
                "Presupuesto": df["Cumple presupuesto"].map(lambda v: pill("Dentro", "ok") if v else pill("Excede", "bad")),
            })
        e1, e2 = st.columns(2)
        e1.markdown("**10 escenarios más rápidos**")
        e1.markdown(tabla_html(fmt_ext(df_sims.nsmallest(10, "Duración (días)"))), unsafe_allow_html=True)
        e2.markdown("**10 escenarios más lentos**")
        e2.markdown(tabla_html(fmt_ext(df_sims.nlargest(10, "Duración (días)"))), unsafe_allow_html=True)

    st.divider()
    seccion("Exportar la simulación", "El Excel incluye resumen, percentiles, histograma y curva S nativos, y todas las iteraciones.", "archive")
    d1, d2, _ = st.columns([1, 1, 1])
    d1.download_button("Reporte Monte Carlo (Excel)", generar_excel_montecarlo(dias, costos, cfg),
                       f"MonteCarlo_{slug(ss.nombre_proj)}.xlsx", XLSX, icon=":material/table_view:", **ANCHO)
    d2.download_button("Iteraciones en CSV", df_sims.to_csv(index=False).encode("utf-8"),
                       f"MonteCarlo_iteraciones_{slug(ss.nombre_proj)}.csv", "text/csv", icon=":material/download:", **ANCHO)


def pagina_riesgos():
    encabezado("shield", "Matriz de riesgos", "Severidad = probabilidad × impacto, en escala de 1 a 25. Sobre 15 se considera crítico.")
    riesgos = st.session_state.riesgos_list
    df = pd.DataFrame(riesgos)
    if not df.empty:
        df["Severidad"] = df["Probabilidad"].astype(int) * df["Impacto"].astype(int)
    franja([
        {"icono": "list", "etiqueta": "Riesgos registrados", "valor": len(df)},
        {"icono": "alert", "etiqueta": "Altos o críticos", "valor": int((df["Severidad"] >= 12).sum()) if not df.empty else 0,
         "sub": "Severidad 12 o más", "destacado": True},
        {"icono": "gauge", "etiqueta": "Severidad promedio",
         "valor": f"{fmt_num(df['Severidad'].mean() if not df.empty else 0, 1)}", "sub": "de 25"},
    ])

    st.write("")
    c_map, c_form = st.columns([1.35, 1], gap="large")
    with c_map:
        seccion("Mapa de calor", icono="grid")
        z = [[i * j for i in range(1, 6)] for j in range(1, 6)]
        fig = go.Figure(go.Heatmap(
            z=z, x=[1, 2, 3, 4, 5], y=[1, 2, 3, 4, 5], showscale=False, hoverinfo="skip", xgap=3, ygap=3,
            text=z, texttemplate="%{text}", textfont=dict(color=rgba(T["text"], 0.35), size=11),
            colorscale=[[0.0, rgba(T["ok"], 0.22)], [0.35, rgba(T["warn"], 0.25)], [0.6, rgba(T["warn"], 0.4)], [1.0, rgba(T["bad"], 0.45)]]))
        if not df.empty:
            fig.add_trace(go.Scatter(
                x=df["Probabilidad"], y=df["Impacto"], mode="markers+text", text=df["Riesgo"],
                textposition=["top right" if int(x) <= 1 else ("top left" if int(x) >= 5 else "top center") for x in df["Probabilidad"]],
                textfont=dict(color=T["text"], size=11), customdata=df["Severidad"],
                marker=dict(size=14, color=T["text"], line=dict(width=2, color=T["card"])),
                hovertemplate="<b>%{text}</b><br>Probabilidad %{x} | Impacto %{y}<br>Severidad %{customdata}<extra></extra>"))
        estilo_fig(fig, height=430, showlegend=False)
        fig.update_xaxes(title="Probabilidad", range=[0.5, 5.5], dtick=1, showgrid=False)
        fig.update_yaxes(title="Impacto", range=[0.5, 5.5], dtick=1, showgrid=False)
        mostrar_fig(fig, key="riesgos_mapa")
    with c_form:
        seccion("Registrar riesgo", icono="flag")
        with st.form("form_add_riesgo", clear_on_submit=True):
            nombre = st.text_input("Descripción del riesgo", placeholder="Ej: Retraso en entregables")
            p = st.slider("Probabilidad", 1, 5, 3)
            i = st.slider("Impacto", 1, 5, 3)
            if st.form_submit_button("Agregar riesgo", icon=":material/add:", **ANCHO):
                if nombre.strip():
                    riesgos.append({"Riesgo": nombre.strip(), "Probabilidad": int(p), "Impacto": int(i)})
                    st.rerun()
                else:
                    st.error("Escribe una descripción para el riesgo.")

    st.divider()
    seccion("Inventario", icono="list")
    if not riesgos:
        st.info("No hay riesgos. Registra el primero con el formulario.")
    else:
        anchos = [3.2, 1.1, 1.1, 0.8, 1.1, 0.45]
        cab = st.columns(anchos)
        for c, t in zip(cab, ["Riesgo", "Probabilidad", "Impacto", "Severidad", "Nivel", ""]):
            c.caption(t)
        borrar = None
        v = version_lista("riesgos_list")
        for idx, r in enumerate(riesgos):
            rc = st.columns(anchos, vertical_alignment="center")
            r["Riesgo"] = rc[0].text_input(f"rt_{idx}", value=r["Riesgo"], key=f"r_name_{v}_{idx}", label_visibility="collapsed")
            r["Probabilidad"] = rc[1].number_input(f"rp_{idx}", value=int(r["Probabilidad"]), min_value=1, max_value=5,
                                                   key=f"r_p_{v}_{idx}", label_visibility="collapsed")
            r["Impacto"] = rc[2].number_input(f"ri_{idx}", value=int(r["Impacto"]), min_value=1, max_value=5,
                                              key=f"r_i_{v}_{idx}", label_visibility="collapsed")
            sev = int(r["Probabilidad"]) * int(r["Impacto"])
            rc[3].markdown(f"**{sev}**")
            rc[4].markdown(pill_riesgo(sev), unsafe_allow_html=True)
            if rc[5].button("", icon=":material/delete:", key=f"del_r_{v}_{idx}", help="Eliminar riesgo"):
                borrar = idx
        if borrar is not None:
            eliminar_fila("riesgos_list", borrar, "r_name_", "r_p_", "r_i_", "del_r_")
            st.rerun()

    st.divider()
    seccion("Exportar", icono="archive")
    d1, _ = st.columns([1, 2])
    d1.download_button("Matriz de riesgos en Excel", generar_excel_riesgos(),
                       f"Matriz_Riesgos_{slug(st.session_state.nombre_proj)}.xlsx", XLSX, icon=":material/table_view:", **ANCHO)


def pagina_reporte():
    encabezado("doc", "Reporte Memento Mori",
               "Documento Word con el análisis crítico de todos los componentes del proyecto, "
               "incluidos los resultados reales de la simulación Monte Carlo.")
    c1, c2 = st.columns([1.3, 1], gap="large")
    with c1:
        seccion("Datos de portada", icono="users")
        a, b = st.columns(2)
        a.text_input("Nombre del proyecto", key="nombre_proj")
        b.text_input("Integrantes / responsables", key="integrantes")
        st.write("")
        seccion("Contenido del reporte", icono="list")
        capitulos = [
            ("Introducción y alcance", "Viabilidad, alineación estratégica y prevención de scope creep."),
            ("Requisitos", f"{len(st.session_state.rf_list)} RF y {len(st.session_state.rnf_list)} RNF, con análisis de vulnerabilidad."),
            ("Arquitectura y flujos", "Acoplamiento de componentes y puntos de fallo."),
            ("Gestión Kanban", "Trabajo en proceso y cuellos de botella."),
            ("EDT", "Actividades críticas por costo y horas."),
            ("Simulación Monte Carlo", "Niveles de confianza, reservas y veredicto." if MC_ACTUAL else "Parámetros inválidos: revisa el módulo."),
            ("Matriz de riesgos", "Planificación anticipada ante fallos y mitigación."),
            ("Conclusiones", "Recomendaciones finales."),
        ]
        html('<div class="mx-panel">' + "".join(
            f'<div class="mx-todo"><span class="pill neu">{n}</span><span><b>{t}</b><br>'
            f'<span style="color:{T["sub"]};font-size:.88rem">{d}</span></span></div>'
            for n, (t, d) in enumerate(capitulos, 1)) + "</div>")
    with c2:
        seccion("Descargar", icono="archive")
        st.download_button("Descargar reporte (.docx)", generar_word_mementomori(),
                           f"Reporte_MementoMori_{slug(st.session_state.nombre_proj)}.docx", DOCX,
                           icon=":material/description:", type="primary", on_click=marcar_descarga, args=("reporte",), **ANCHO)
        if MC_ACTUAL is None:
            st.warning("La simulación tiene parámetros inválidos; el capítulo 6 lo indicará.")
        st.write("")
        html(panel("Datos que se incluirán", [
            ("Proyecto", st.session_state.nombre_proj),
            ("Equipo", st.session_state.integrantes),
            ("Tareas", f"{len(st.session_state.kanban_tasks)}"),
            ("Paquetes EDT", f"{len(st.session_state.edt_list)}"),
            ("Riesgos", f"{len(st.session_state.riesgos_list)}"),
            ("Plazo comprometido", f"{fmt_num(MC_ACTUAL['met']['pt'], 1)} días" if MC_ACTUAL else "—"),
        ], "info"))


def pagina_descargas():
    encabezado("archive", "Descarga completa", "Todos los entregables del proyecto en un único archivo ZIP.")
    nombre_s = slug(st.session_state.nombre_proj)
    archivos = [
        (f"Reporte_MementoMori_{nombre_s}.docx", "Reporte crítico integral", "doc"),
        (f"ERS_{nombre_s}.docx", "Especificación de requisitos", "layers"),
        (f"Matrices_{nombre_s}.xlsx", "Requisitos, Kanban, EDT, riesgos y resumen Monte Carlo", "grid"),
        (f"MonteCarlo_{nombre_s}.xlsx", "Simulación completa con gráficos" if MC_ACTUAL else "No disponible (parámetros inválidos)", "curve"),
        ("Kanban_Tareas.csv / EDT_WBS.csv", "Datos en texto plano", "list"),
        ("Mermaid/*.mmd", "Código de los diagramas", "nodes"),
    ]
    c1, c2 = st.columns([1.4, 1], gap="large")
    with c1:
        html('<div class="mx-panel">' + "".join(
            f'<div class="mx-todo">{ico(i, 18, T["accent"])}<span><b>{n}</b><br>'
            f'<span style="color:{T["sub"]};font-size:.88rem">{d}</span></span></div>' for n, d, i in archivos) + "</div>")
    with c2:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr(f"Reporte_MementoMori_{nombre_s}.docx", generar_word_mementomori())
            zf.writestr(f"ERS_{nombre_s}.docx", generar_word_ers())
            zf.writestr(f"Matrices_{nombre_s}.xlsx", generar_excel_estilizado())
            if MC_ACTUAL:
                zf.writestr(f"MonteCarlo_{nombre_s}.xlsx",
                            generar_excel_montecarlo(MC_ACTUAL["res"]["dias"], MC_ACTUAL["res"]["costos"], MC_ACTUAL["cfg"]))
            zf.writestr("Kanban_Tareas.csv", pd.DataFrame(st.session_state.kanban_tasks).to_csv(index=False).encode("utf-8"))
            zf.writestr("EDT_WBS.csv", pd.DataFrame(st.session_state.edt_list).to_csv(index=False).encode("utf-8"))
            zf.writestr("Mermaid/Modelado.mmd", st.session_state.modelo_mermaid)
            zf.writestr("Mermaid/Flujo.mmd", st.session_state.flujo_codigo)
        st.download_button("Descargar paquete completo (.zip)", buf.getvalue(), f"Reporte_Full_{nombre_s}.zip",
                           "application/zip", icon=":material/folder_zip:", type="primary",
                           on_click=marcar_descarga, args=("descargas",), **ANCHO)
        st.caption(f"Tamaño aproximado: {fmt_num(len(buf.getvalue()) / 1024, 0)} KB")


# =====================================================================
# 11. NAVEGACIÓN
# =====================================================================
def _envolver(funcion, clave):
    """Registra la visita, dibuja la página y agrega los enlaces de paso anterior / siguiente."""
    def _pagina():
        marcar_visita(clave)
        funcion()
        pie_navegacion(clave)
    _pagina.__name__ = f"pagina_{clave}"
    return _pagina


ICONOS_PAGINA = {
    "requisitos": ":material/checklist:", "modelado": ":material/account_tree:", "flujos": ":material/route:",
    "kanban": ":material/view_kanban:", "edt": ":material/lan:", "montecarlo": ":material/query_stats:",
    "riesgos": ":material/gpp_maybe:", "reporte": ":material/description:", "descargas": ":material/folder_zip:",
}
FUNCIONES_PAGINA = {
    "requisitos": pagina_requisitos, "modelado": pagina_modelado, "flujos": pagina_flujos, "kanban": pagina_gestion,
    "edt": pagina_edt, "montecarlo": pagina_montecarlo, "riesgos": pagina_riesgos, "reporte": pagina_reporte,
    "descargas": pagina_descargas,
}

PG_INICIO = st.Page(pagina_inicio, title="Inicio", icon=":material/home:", url_path="inicio", default=True)
PAGINAS = {
    clave: st.Page(_envolver(FUNCIONES_PAGINA[clave], clave), title=f"{n}. {titulo}", icon=ICONOS_PAGINA[clave], url_path=clave)
    for n, (clave, titulo, *_resto) in enumerate(PASOS, 1)
}
PG_REQ, PG_MODELADO, PG_FLUJOS = PAGINAS["requisitos"], PAGINAS["modelado"], PAGINAS["flujos"]
PG_KANBAN, PG_EDT, PG_MC, PG_RIESGOS = PAGINAS["kanban"], PAGINAS["edt"], PAGINAS["montecarlo"], PAGINAS["riesgos"]
PG_REPORTE, PG_DESCARGAS = PAGINAS["reporte"], PAGINAS["descargas"]

navegacion = st.navigation({
    "Bienvenida": [PG_INICIO],
    "Definición (pasos 1 a 3)": [PG_REQ, PG_MODELADO, PG_FLUJOS],
    "Planificación (pasos 4 a 7)": [PG_KANBAN, PG_EDT, PG_MC, PG_RIESGOS],
    "Entregables (pasos 8 y 9)": [PG_REPORTE, PG_DESCARGAS],
})

# Si el tema nativo cambió, se recarga la página actual una vez para que se aplique de inmediato
if aplicar_tema_nativo(T) and st.session_state.get("_tema_sincronizado") != st.session_state.tema:
    st.session_state._tema_sincronizado = st.session_state.tema
    st.switch_page(navegacion)
st.session_state._tema_sincronizado = st.session_state.tema

def alternar_modo_oscuro():
    st.session_state.tema = "claro" if st.session_state.tema in ("oscuro", "informatico") else "oscuro"


# Botón colgante (fijo arriba a la derecha): luna para pasar a oscuro, sol para volver a claro
es_oscuro = st.session_state.tema in ("oscuro", "informatico")
st.button("", key="btn_luna", icon=":material/light_mode:" if es_oscuro else ":material/dark_mode:",
          on_click=alternar_modo_oscuro, help="Volver al modo claro" if es_oscuro else "Cambiar a modo oscuro")

for _clave, _pg in PAGINAS.items():
    if _pg.url_path == navegacion.url_path:
        marcar_visita(_clave)
_estados, _hechos, _sig = avance_proyecto()
with st.sidebar:
    html(f'<div class="mx-brand">{ico("hub", 22)}<div>MatrixDev<small>Gestión de proyectos de software</small></div></div>')
    html(f'<div class="mx-sideprog"><div class="top"><span>Avance</span><b>{_hechos} de {len(PASOS)} pasos</b></div>'
         f'<div class="mx-bar" style="--c:{T["accent"]};--w:{_hechos / len(PASOS) * 100:.1f}%;--m:101%"><div class="tr"><div class="fl"></div></div></div>'
         + (f'<div class="nx">Siguiente: {_sig[0]}. {_sig[1][1]}</div>' if _sig else '<div class="nx">Todos los pasos completos</div>')
         + '</div>')
    st.segmented_control("Tema", list(TEMAS.keys()), key="tema", format_func=lambda k: TEMAS[k]["icono"],
                         help="Claro (celeste y blanco), Oscuro (negro y blanco) o Informático (negro y rojo neón).")
    st.caption(f"Tema activo: {T['nombre']}")
    if estado_paso("montecarlo")[0]:
        st.caption(f"Plazo P{MC_ACTUAL['cfg']['perc']}: {fmt_num(MC_ACTUAL['met']['pt'], 1)} días  |  "
                   f"Éxito conjunto: {fmt_num(MC_ACTUAL['met']['prob_ambos'], 1)}%")

navegacion.run()
